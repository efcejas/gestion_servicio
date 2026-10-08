from decimal import Decimal
from collections import defaultdict
import json
import re
import unicodedata
from uuid import UUID, uuid4

from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone

from .models import (
    GrupoTarifario,
    HistorialAjusteCantidadDopplerMMII,
    HistorialRevisionAuditoriaDopplerMMII,
    RevisionAuditoriaDopplerMMII,
    RegistroEstudio,
    RegistroEstudiosPorMedico,
    RevisionCruceEgesRegistro,
    SesionContable,
)
from control_guardias.models import Feriado


TIPO_SIN_REGISTRO_ESTUDIO = 'sin_registro_estudio'
TIPO_MONTO_CERO_CON_ESTUDIOS = 'monto_cero_con_estudios'
TIPO_GUARDIA_MONTO_INVALIDO = 'guardia_monto_invalido'
TIPO_SESION_VACIA = 'sesion_vacia'
TIPO_SIN_GRUPO_CON_FALLBACK = 'sin_grupo_con_fallback'
TIPO_SIN_PRECIO_RESOLUBLE = 'sin_precio_resoluble'
TIPO_SIN_TARIFA_VIGENTE_GRUPO = 'sin_tarifa_vigente_grupo'
TIPO_CONTEXTUAL_SIN_GRUPO = 'contextual_sin_grupo'
TIPO_CONTEXTUAL_SIN_TARIFA = 'contextual_sin_tarifa'


# Auditoría administrativa residentes ECO (PR2)
ROLES_AUDITORIA_RESIDENCIA_ECO = {
    'medico_residente',
    'jefe_residentes',
    'instructor_residentes',
}
TIPOS_AUDITORIA_RESIDENCIA_ECO = {'ECO', 'DOP', 'ECOCAR'}

AUDIT_EXTRA_MENSUAL_AMARILLA = 35
AUDIT_EXTRA_MENSUAL_ROJA = 50
AUDIT_PROP_EXTRA_AMARILLA = 0.35
AUDIT_PROP_EXTRA_ROJA = 0.50
AUDIT_NOCTURNOS_AMARILLA = 6
AUDIT_NOCTURNOS_ROJA = 10
AUDIT_FINDE_FERIADO_AMARILLA = 8
AUDIT_FINDE_FERIADO_ROJA = 15
AUDIT_MAX_ECO_DIA_AMARILLA = 14
AUDIT_MAX_ECO_DIA_ROJA = 20
AUDIT_HORA_NOCTURNA_DESDE = 22
AUDIT_HORA_NOCTURNA_HASTA = 6
AUDIT_HORA_POST_17 = 17


def _normalizar_texto_doppler(valor):
    texto = ''.join(
        caracter
        for caracter in unicodedata.normalize('NFKD', str(valor or '').upper())
        if not unicodedata.combining(caracter)
    )
    return re.sub(r'[^A-Z0-9]+', ' ', texto).strip()


def _clasificar_doppler_mmii(estudio):
    nombre = _normalizar_texto_doppler(estudio.nombre)
    codigo = _normalizar_texto_doppler(estudio.codigo)
    grupo = getattr(estudio, 'grupo_tarifario', None)
    modalidad_grupo = (getattr(grupo, 'modalidad', '') or '').upper()
    texto = f'{nombre} {codigo}'

    es_doppler = (
        (estudio.tipo or '').upper() == 'DOP'
        or modalidad_grupo == 'DOP'
        or 'DOPPLER' in texto
        or 'ECODOPPLER' in texto
    )
    es_mmii = bool(re.search(r'\bMMII\b|\bMM\s+INFERIORES\b|MIEMBROS INFERIORES', texto))
    if not es_doppler or not es_mmii:
        return None

    es_arterial = 'ARTERIAL' in nombre
    es_venoso = 'VENOSO' in nombre or 'VENOSA' in nombre
    if es_arterial and es_venoso:
        return {'tipo': 'ARTERIAL_Y_VENOSO_MMII', 'cantidad_esperada': 2}
    if es_arterial:
        return {'tipo': 'ARTERIAL_MMII', 'cantidad_esperada': 1}
    if es_venoso:
        return {'tipo': 'VENOSO_MMII', 'cantidad_esperada': 1}
    return {'tipo': 'DOPPLER_MMII_REVISAR', 'cantidad_esperada': None}


def auditar_cantidad_doppler_mmii(*, fecha_desde, fecha_hasta, medico_id=None):
    """Devuelve diferencias potenciales sin modificar registros ni montos."""
    registros = (
        RegistroEstudiosPorMedico.objects
        .filter(
            anulado=False,
            fecha_del_informe__range=(fecha_desde, fecha_hasta),
        )
        .select_related('medico')
        .prefetch_related(
            'registroestudio_set__estudio__grupo_tarifario',
        )
        .order_by('fecha_del_informe', 'medico__last_name', 'medico__first_name', 'pk')
    )
    if medico_id:
        registros = registros.filter(medico_id=medico_id)

    resultados = []
    posibles_duplicados = defaultdict(list)
    estudios_por_relacion = {}
    resultados_por_registro = {}
    for registro in registros:
        dopplers = []
        regiones_ajuste = 0
        for relacion in registro.registroestudio_set.all():
            regla = _clasificar_doppler_mmii(relacion.estudio)
            if not regla:
                continue

            cantidad_esperada = regla['cantidad_esperada']
            if cantidad_esperada is None:
                regiones_esperadas = None
            else:
                regiones_esperadas = (
                    relacion.estudio.conteo_regiones_default * cantidad_esperada
                )
                regiones_ajuste += (
                    relacion.estudio.conteo_regiones_default
                    * (cantidad_esperada - relacion.cantidad)
                )

            dopplers.append({
                'registro_estudio_id': relacion.pk,
                'estudio_id': relacion.estudio_id,
                'estudio': relacion.estudio.nombre,
                'codigo_estudio': relacion.estudio.codigo or '',
                'clasificacion': regla['tipo'],
                'cantidad_declarada': relacion.cantidad,
                'cantidad_esperada': cantidad_esperada,
                'diferencia_cantidad': (
                    relacion.cantidad - cantidad_esperada
                    if cantidad_esperada is not None
                    else None
                ),
                'regiones_declaradas': (
                    relacion.estudio.conteo_regiones_default * relacion.cantidad
                ),
                'regiones_esperadas': regiones_esperadas,
                'contexto': relacion.contexto,
                'requiere_revision_manual': cantidad_esperada is None,
                'posible_duplicado': False,
                'cantidad_en_grupo_duplicado': 1,
                'otros_registros_posible_duplicado': [],
            })
            estudios_por_relacion[relacion.pk] = dopplers[-1]
            dni_normalizado = re.sub(r'\D', '', str(registro.dni_paciente or ''))
            if dni_normalizado and cantidad_esperada is not None:
                posibles_duplicados[(
                    registro.medico_id,
                    dni_normalizado,
                    registro.fecha_del_informe,
                    regla['tipo'],
                )].append({
                    'registro_id': registro.pk,
                    'registro_estudio_id': relacion.pk,
                    'estudio': relacion.estudio.nombre,
                })

        if not dopplers:
            continue

        cantidades_revision_manual = any(
            item['requiere_revision_manual'] for item in dopplers
        )
        cantidad_practicas_declarada = sum(
            item['cantidad_declarada'] for item in dopplers
        )
        cantidad_practicas_esperada = (
            None
            if cantidades_revision_manual
            else sum(item['cantidad_esperada'] for item in dopplers)
        )
        diferencia_cantidad = sum(
            item['cantidad_declarada'] - item['cantidad_esperada']
            for item in dopplers
            if item['cantidad_esperada'] is not None
        )
        resultado_registro = {
            'registro_id': registro.pk,
            'sesion_id': registro.sesion_contable_id,
            'fecha': registro.fecha_del_informe,
            'medico_id': registro.medico_id,
            'medico': registro.medico.get_full_name() or registro.medico.username,
            'rol': registro.medico.rol,
            'paciente': f'{registro.apellido_paciente}, {registro.nombre_paciente}',
            'dni': registro.dni_paciente,
            'tipo_obra_social': registro.tipo_obra_social,
            'horario': registro.horario,
            'cantidad_practicas_declarada': cantidad_practicas_declarada,
            'cantidad_practicas_esperada': cantidad_practicas_esperada,
            'cantidad_regiones_declarada': registro.cantidad_regiones,
            'cantidad_regiones_esperada': max(
                registro.cantidad_regiones + regiones_ajuste,
                0,
            ),
            'monto_registrado': registro.monto_calculado,
            'diferencia_cantidad': diferencia_cantidad,
            'requiere_revision_manual': cantidades_revision_manual,
            'tratamiento_economico': (
                'INFORMATIVO_SIN_DEBITO'
                if registro.medico.rol == 'medico_residente'
                else 'IMPACTO_POTENCIAL_NO_APLICADO'
            ),
            'estudios': dopplers,
        }
        resultados.append(resultado_registro)
        resultados_por_registro[registro.pk] = resultado_registro

    grupos_duplicados = 0
    for coincidencias in posibles_duplicados.values():
        if len(coincidencias) < 2:
            continue
        grupos_duplicados += 1
        ids_registro = sorted({item['registro_id'] for item in coincidencias})
        for coincidencia in coincidencias:
            estudio = estudios_por_relacion[coincidencia['registro_estudio_id']]
            resultado = resultados_por_registro[coincidencia['registro_id']]
            estudio['posible_duplicado'] = True
            estudio['cantidad_en_grupo_duplicado'] = len(coincidencias)
            estudio['otros_registros_posible_duplicado'] = [
                pk for pk in ids_registro if pk != resultado['registro_id']
            ]
            estudio['requiere_revision_manual'] = True
            resultado['requiere_revision_manual'] = True

    return {
        'fecha_desde': fecha_desde,
        'fecha_hasta': fecha_hasta,
        'total_registros': len(resultados),
        'total_con_diferencia': sum(
            1 for item in resultados
            if (
                item['diferencia_cantidad'] != 0
                or item['requiere_revision_manual']
                or any(estudio.get('posible_duplicado') for estudio in item['estudios'])
            )
        ),
        'grupos_posible_duplicado': grupos_duplicados,
        'resultados': resultados,
    }


@transaction.atomic
def crear_casos_auditoria_doppler_mmii(auditoria, usuario):
    """Persiste candidatos una sola vez, conservando un snapshot de los datos fuente."""
    candidatos = []
    registro_ids = []
    fuentes = {
        registro.pk: fuente_calculo_doppler_mmii(registro)
        for registro in RegistroEstudiosPorMedico.objects.filter(
            pk__in=[resultado['registro_id'] for resultado in auditoria.get('resultados', [])],
        ).select_related('medico').prefetch_related('registroestudio_set')
    }
    for resultado in auditoria.get('resultados', []):
        registro_ids.append(resultado['registro_id'])
        for estudio in resultado['estudios']:
            requiere_revision = (
                estudio['requiere_revision_manual']
                or estudio['posible_duplicado']
                or estudio['diferencia_cantidad'] not in (None, 0)
            )
            if not requiere_revision:
                continue

            snapshot = {
                'fuente_calculo': fuentes[resultado['registro_id']],
                'regla_version': RevisionAuditoriaDopplerMMII.VERSION_REGLA_V1,
                'registro_id': resultado['registro_id'],
                'registro_estudio_id': estudio['registro_estudio_id'],
                'profesional_id': resultado['medico_id'],
                'profesional': resultado['medico'],
                'rol': resultado['rol'],
                'paciente': resultado['paciente'],
                'dni': resultado['dni'],
                'fecha_informe': resultado['fecha'].isoformat() if resultado['fecha'] else None,
                'tipo_obra_social': resultado.get('tipo_obra_social'),
                'horario': resultado.get('horario'),
                'monto_registrado': str(resultado['monto_registrado']),
                'regiones_registro_declaradas': resultado['cantidad_regiones_declarada'],
                'estudio_id': estudio['estudio_id'],
                'estudio': estudio['estudio'],
                'codigo_estudio': estudio['codigo_estudio'],
                'clasificacion': estudio['clasificacion'],
                'contexto': estudio['contexto'],
                'cantidad_declarada': estudio['cantidad_declarada'],
                'cantidad_esperada': estudio['cantidad_esperada'],
                'regiones_declaradas': estudio['regiones_declaradas'],
                'regiones_esperadas': estudio['regiones_esperadas'],
                'posible_duplicado': estudio['posible_duplicado'],
                'otros_registros_posible_duplicado': estudio['otros_registros_posible_duplicado'],
            }
            candidatos.append(RevisionAuditoriaDopplerMMII(
                registro_id=resultado['registro_id'],
                registro_estudio_id=estudio['registro_estudio_id'],
                registro_estudio_id_origen=estudio['registro_estudio_id'],
                version_regla=RevisionAuditoriaDopplerMMII.VERSION_REGLA_V1,
                datos_originales_json=snapshot,
                creado_por=usuario,
            ))

    if not candidatos:
        return {'detectados': 0, 'creados': 0, 'existentes': 0}

    claves_candidatas = {
        (caso.registro_id, caso.registro_estudio_id_origen)
        for caso in candidatos
    }
    casos_existentes = set(
        RevisionAuditoriaDopplerMMII.objects
        .filter(
            registro_id__in=registro_ids,
            version_regla=RevisionAuditoriaDopplerMMII.VERSION_REGLA_V1,
        )
        .values_list('registro_id', 'registro_estudio_id_origen')
    )
    existentes = len(claves_candidatas & casos_existentes)
    RevisionAuditoriaDopplerMMII.objects.bulk_create(candidatos, ignore_conflicts=True)
    return {
        'detectados': len(claves_candidatas),
        'creados': len(claves_candidatas) - existentes,
        'existentes': existentes,
    }


def adjuntar_revisiones_auditoria_doppler_mmii(auditoria):
    """Adjunta la revision vigente a cada linea del informe."""
    estudios = [
        estudio
        for resultado in auditoria.get('resultados', [])
        for estudio in resultado['estudios']
    ]
    linea_ids = [estudio['registro_estudio_id'] for estudio in estudios]
    revisiones = (
        RevisionAuditoriaDopplerMMII.objects
        .filter(
            registro_estudio_id_origen__in=linea_ids,
            version_regla=RevisionAuditoriaDopplerMMII.VERSION_REGLA_V1,
        )
        .select_related('revisado_por', 'creado_por')
        .prefetch_related('historial__revisado_por')
        .order_by('registro_estudio_id_origen', '-fecha_actualizacion')
    ) if linea_ids else []
    revision_por_linea = {}
    for revision in revisiones:
        revision_por_linea.setdefault(revision.registro_estudio_id_origen, revision)

    for estudio in estudios:
        estudio['revision_auditoria'] = revision_por_linea.get(
            estudio['registro_estudio_id'],
        )
        caso = estudio['revision_auditoria']
        estudio['permite_lote'] = bool(
            caso and caso_doppler_permite_lote(caso)
            and not estudio['posible_duplicado']
            and estudio['cantidad_declarada'] == 2 and estudio['cantidad_esperada'] == 1
        )
    return auditoria


@transaction.atomic
def resolver_caso_auditoria_doppler_mmii(*, caso_id, decision, evidencias, observacion, usuario):
    """Registra la decision actual y agrega un evento inmutable de historial."""
    if decision not in {
        RevisionAuditoriaDopplerMMII.ESTADO_CONFIRMADO,
        RevisionAuditoriaDopplerMMII.ESTADO_DESCARTADO,
        RevisionAuditoriaDopplerMMII.ESTADO_REQUIERE_EVIDENCIA,
    }:
        raise ValidationError('La decision de auditoria no es valida.')
    if not observacion.strip() or len(observacion.strip()) > 2000:
        raise ValidationError('Indica una observacion de hasta 2000 caracteres.')
    if decision == RevisionAuditoriaDopplerMMII.ESTADO_CONFIRMADO and not (
        evidencias.get('orden_medica_verificada')
        or evidencias.get('visualmedical_verificado')
    ):
        raise ValidationError(
            'Para confirmar la diferencia debes verificar la orden o VisualMedical.'
        )

    registro_id = RevisionAuditoriaDopplerMMII.objects.values_list(
        'registro_id', flat=True,
    ).get(pk=caso_id)
    registro = RegistroEstudiosPorMedico.objects.select_for_update().get(pk=registro_id)
    caso = RevisionAuditoriaDopplerMMII.objects.select_for_update().get(pk=caso_id)
    estado_anterior = caso.estado
    caso.estado = decision
    caso.orden_medica_verificada = evidencias.get('orden_medica_verificada', False)
    caso.eges_verificado = evidencias.get('eges_verificado', False)
    caso.visualmedical_verificado = evidencias.get('visualmedical_verificado', False)
    caso.netterm_verificado = evidencias.get('netterm_verificado', False)
    caso.observacion = observacion.strip()
    caso.revisado_por = usuario
    caso.fecha_revision = timezone.now()
    caso.save(update_fields=[
        'estado',
        'orden_medica_verificada',
        'eges_verificado',
        'visualmedical_verificado',
        'netterm_verificado',
        'observacion',
        'revisado_por',
        'fecha_revision',
        'fecha_actualizacion',
    ])
    caso.estimacion_json = estimar_registro_auditoria_doppler_mmii(registro)
    caso.save(update_fields=['estimacion_json'])
    HistorialRevisionAuditoriaDopplerMMII.objects.create(
        revision=caso,
        estado_anterior=estado_anterior,
        estado_nuevo=decision,
        estimacion_json=caso.estimacion_json,
        orden_medica_verificada=caso.orden_medica_verificada,
        eges_verificado=caso.eges_verificado,
        visualmedical_verificado=caso.visualmedical_verificado,
        netterm_verificado=caso.netterm_verificado,
        observacion=caso.observacion,
        revisado_por=usuario,
    )
    return caso


def fuente_calculo_doppler_mmii(registro):
    return {
        'medico_id': registro.medico_id,
        'rol': registro.medico.rol,
        'remoto': registro.medico.trabaja_remoto,
        'fecha': registro.fecha_del_informe.isoformat(),
        'horario': registro.horario,
        'obra_social': registro.tipo_obra_social,
        'monto': str(registro.monto_calculado),
        'regiones': registro.cantidad_regiones,
        'internado': registro.paciente_internado,
        'solicitud': registro.fecha_hora_solicitud.isoformat() if registro.fecha_hora_solicitud else None,
        'informe': registro.fecha_hora_informe.isoformat() if registro.fecha_hora_informe else None,
        'anulado': registro.anulado,
        'lineas': sorted([
            [rel.pk, rel.estudio_id, rel.cantidad, rel.contexto]
            for rel in registro.registroestudio_set.all()
        ]),
    }


def estimar_registro_auditoria_doppler_mmii(registro):
    """Simula solo cantidades confirmadas, sin escribir prestaciones ni montos."""
    casos = list(registro.revisiones_auditoria_doppler_mmii.filter(
        version_regla=RevisionAuditoriaDopplerMMII.VERSION_REGLA_V1,
    ).order_by('fecha_deteccion', 'pk'))
    original = casos[0].datos_originales_json.get('monto_registrado') if casos else None
    resultado = {
        'monto_original': original,
        'monto_estimado': None,
        'diferencia_estimada': None,
        'motivo': 'Sin cantidades confirmadas.',
        'informativo_residente': registro.medico.rol == 'medico_residente',
        'casos_confirmados': [caso.pk for caso in casos if caso.estado == 'CONFIRMADO'],
        'parcial': any(caso.estado in {'PENDIENTE', 'REQUIERE_EVIDENCIA'} for caso in casos),
    }
    confirmados = [caso for caso in casos if caso.estado == 'CONFIRMADO']
    if not confirmados:
        return resultado
    relaciones = {rel.pk: rel for rel in registro.registroestudio_set.select_related('estudio').all()}
    fuente_actual = fuente_calculo_doppler_mmii(registro)
    cantidades = {}
    for caso in casos:
        fuente = caso.datos_originales_json
        relacion = relaciones.get(caso.registro_estudio_id_origen)
        if (
            registro.anulado
            or (fuente.get('fuente_calculo') and fuente['fuente_calculo'] != fuente_actual)
            or not relacion
            or relacion.estudio_id != fuente.get('estudio_id')
            or relacion.cantidad != fuente.get('cantidad_declarada')
            or relacion.contexto != fuente.get('contexto')
            or registro.fecha_del_informe.isoformat() != fuente.get('fecha_informe')
            or registro.tipo_obra_social != fuente.get('tipo_obra_social')
            or registro.horario != fuente.get('horario')
            or registro.medico_id != fuente.get('profesional_id')
            or registro.medico.rol != fuente.get('rol')
            or str(registro.monto_calculado) != fuente.get('monto_registrado')
        ):
            resultado['motivo'] = 'Los datos actuales no coinciden con el snapshot; requiere revision.'
            return resultado
        if caso.estado != 'CONFIRMADO':
            continue
        esperada = fuente.get('cantidad_esperada')
        if fuente.get('posible_duplicado') or not isinstance(esperada, int) or esperada < 1:
            resultado['motivo'] = 'Duplicado o cantidad manual: requiere una resolucion individual de prestaciones.'
            return resultado
        cantidades[relacion.pk] = esperada

    monto_original = Decimal(original)
    if any(rel.estudio.precio_para_os(
        registro.tipo_obra_social, fecha=registro.fecha_del_informe, contexto=rel.contexto,
    ) <= 0 for rel in relaciones.values()):
        resultado['motivo'] = 'Falta una tarifa positiva para alguna practica; estimacion pendiente.'
        return resultado
    base = registro.calcular_monto().quantize(Decimal('0.01'))
    if base != monto_original or base <= 0:
        resultado['motivo'] = 'La tarifa historica actual no reproduce el monto original; estimacion pendiente.'
        return resultado
    monto_estimado = registro.calcular_monto(cantidades_auditoria=cantidades).quantize(Decimal('0.01'))
    resultado.update({
        'monto_estimado': str(monto_estimado),
        'diferencia_estimada': str(monto_original - monto_estimado),
        'motivo': 'Simulacion de cantidades confirmadas; no es un debito aplicado.',
        'cantidades_simuladas': {str(linea_id): cantidad for linea_id, cantidad in cantidades.items()},
        'fecha_informe': registro.fecha_del_informe.isoformat(),
        'horario': registro.horario,
        'tipo_obra_social': registro.tipo_obra_social,
        'fuente_calculo': fuente_actual,
    })
    return resultado


def caso_doppler_permite_lote(caso):
    fuente = caso.datos_originales_json
    return (
        caso.estado in {'PENDIENTE', 'REQUIERE_EVIDENCIA'}
        and caso.version_regla == RevisionAuditoriaDopplerMMII.VERSION_REGLA_V1
        and not fuente.get('posible_duplicado')
        and fuente.get('cantidad_esperada') == 1
        and fuente.get('cantidad_declarada') == 2
    )


def _validar_caso_aplicacion_cantidad_doppler(caso, registro, relacion):
    fuente = caso.datos_originales_json
    errores = []
    if caso.estado != RevisionAuditoriaDopplerMMII.ESTADO_CONFIRMADO:
        errores.append('El caso no esta confirmado.')
    if not (caso.orden_medica_verificada or caso.visualmedical_verificado):
        errores.append('La confirmacion no conserva orden o VisualMedical verificada.')
    if registro.anulado or registro.medico.rol == 'medico_residente':
        errores.append('El registro esta anulado o es informativo para residente.')
    if (
        not relacion or relacion.registro_id != registro.pk
        or relacion.estudio_id != fuente.get('estudio_id')
        or relacion.pk != caso.registro_estudio_id_origen
        or relacion.cantidad != fuente.get('cantidad_declarada')
        or relacion.contexto != fuente.get('contexto')
    ):
        errores.append('La linea o cantidad declarada cambio desde el snapshot.')
    esperada = fuente.get('cantidad_esperada')
    if fuente.get('posible_duplicado') or esperada != 1 or fuente.get('cantidad_declarada') != 2:
        errores.append('El caso no es una correccion bilateral elegible 2 -> 1.')
    if relacion and relacion.cantidad_liquidable is not None:
        eventos = HistorialAjusteCantidadDopplerMMII.objects.filter(
            revision=caso, registro_estudio_id_origen=relacion.pk,
        ).order_by('-fecha_evento', '-pk')
        ultimo = eventos.first()
        if not ultimo or ultimo.accion != HistorialAjusteCantidadDopplerMMII.ACCION_REVERTIR:
            errores.append('La cantidad liquidable ya tiene un ajuste activo.')
    if fuente.get('fuente_calculo'):
        actual = fuente_calculo_doppler_mmii(registro)
        for campo in ('medico_id', 'rol', 'remoto', 'fecha', 'horario', 'obra_social',
                      'regiones', 'internado', 'solicitud', 'informe', 'anulado', 'lineas'):
            if actual[campo] != fuente['fuente_calculo'].get(campo):
                errores.append(f'El dato fuente {campo} cambio desde la deteccion.')
                break
        if actual['monto'] != fuente['fuente_calculo'].get('monto'):
            reversiones = HistorialAjusteCantidadDopplerMMII.objects.filter(
                accion=HistorialAjusteCantidadDopplerMMII.ACCION_REVERTIR,
                evento_origen__isnull=False,
                evento_origen__registro_id=registro.pk,
            ).values_list('evento_origen_id', flat=True)
            ajuste_activo_monto_actual = HistorialAjusteCantidadDopplerMMII.objects.filter(
                registro=registro,
                accion=HistorialAjusteCantidadDopplerMMII.ACCION_APLICAR,
                monto_registro_nuevo=registro.monto_calculado,
            ).exclude(pk__in=reversiones).exists()
            if not ajuste_activo_monto_actual:
                errores.append('El monto cambio desde el snapshot sin un ajuste Doppler activo que lo respalde.')
    else:
        # Snapshots P1 anteriores a la huella completa mantienen validacion de los campos disponibles.
        if (
            registro.fecha_del_informe.isoformat() != fuente.get('fecha_informe')
            or registro.tipo_obra_social != fuente.get('tipo_obra_social')
            or registro.horario != fuente.get('horario')
            or registro.medico_id != fuente.get('profesional_id')
            or registro.medico.rol != fuente.get('rol')
            or str(registro.monto_calculado) != fuente.get('monto_registrado')
        ):
            errores.append('Los datos actuales ya no coinciden con el snapshot P1.')
    if not registro.sesion_contable_id or registro.sesion_contable.estado not in {'ABIERTA', 'REVISION'}:
        errores.append('La sesion no permite aplicar correcciones economicas.')
    return errores


def preparar_aplicacion_cantidad_doppler_mmii(*, caso_ids, usuario):
    """Devuelve cantidades y montos previstos sin persistir ningun cambio."""
    if not usuario.is_authenticated or not (
        usuario.is_superuser or usuario.rol == 'jefe_servicio'
    ):
        raise ValidationError('No tienes permisos para aplicar ajustes Doppler.')
    ids = list(dict.fromkeys(int(caso_id) for caso_id in caso_ids))
    if not ids or len(ids) > 200 or len(ids) != len(caso_ids):
        raise ValidationError('Selecciona entre 1 y 200 casos distintos.')
    casos = list(RevisionAuditoriaDopplerMMII.objects.filter(
        pk__in=ids,
        version_regla=RevisionAuditoriaDopplerMMII.VERSION_REGLA_V1,
    ).select_related('registro__medico', 'registro__sesion_contable', 'registro_estudio').order_by('registro_id', 'pk'))
    if len(casos) != len(ids):
        raise ValidationError('Uno o mas casos seleccionados ya no existen.')
    por_registro = defaultdict(list)
    errores = []
    for caso in casos:
        registro = caso.registro
        relacion = caso.registro_estudio
        errores_caso = _validar_caso_aplicacion_cantidad_doppler(caso, registro, relacion)
        if errores_caso:
            errores.append({'caso_id': caso.pk, 'registro_id': registro.pk, 'motivos': errores_caso})
        else:
            por_registro[registro.pk].append((caso, relacion))
    if errores:
        raise ValidationError('Seleccion no elegible: ' + json.dumps(errores, ensure_ascii=True))
    resultado = []
    for registros_casos in por_registro.values():
        caso_muestra, _relacion = registros_casos[0]
        registro = caso_muestra.registro
        monto_actual = registro.monto_calculado or Decimal('0.00')
        monto_canonico_actual = registro.calcular_monto().quantize(Decimal('0.01'))
        if monto_canonico_actual != monto_actual:
            raise ValidationError(
                f'El monto actual del registro #{registro.pk} no coincide con el calculo canonico; requiere revision.'
            )
        cantidades = {
            relacion.pk: caso.datos_originales_json['cantidad_esperada']
            for caso, relacion in registros_casos
        }
        monto_nuevo = registro.calcular_monto(cantidades_auditoria=cantidades).quantize(Decimal('0.01'))
        resultado.append({
            'registro': registro,
            'casos': [caso for caso, _relacion in registros_casos],
            'monto_anterior': monto_actual,
            'monto_nuevo': monto_nuevo,
            'diferencia': monto_nuevo - monto_actual,
            'cantidades': cantidades,
        })
    return resultado


@transaction.atomic
def aplicar_lote_cantidad_doppler_mmii(*, caso_ids, usuario, motivo):
    """Aplica la cantidad liquidable confirmada, sin sobrescribir la cantidad declarada."""
    motivo = (motivo or '').strip()
    if not motivo or len(motivo) > 2000:
        raise ValidationError('Indica un fundamento de hasta 2000 caracteres.')
    plan = preparar_aplicacion_cantidad_doppler_mmii(caso_ids=caso_ids, usuario=usuario)
    sesion_ids = sorted({item['registro'].sesion_contable_id for item in plan})
    list(SesionContable.objects.select_for_update().filter(
        pk__in=sesion_ids,
    ).order_by('pk'))
    registro_ids = sorted(item['registro'].pk for item in plan)
    list(RegistroEstudiosPorMedico.objects.select_for_update().filter(
        pk__in=registro_ids,
    ).order_by('pk'))
    caso_por_id = {caso.pk: caso for item in plan for caso in item['casos']}
    casos_bloqueados = list(RevisionAuditoriaDopplerMMII.objects.select_for_update().filter(
        pk__in=caso_por_id,
    ).order_by('registro_id', 'pk'))
    list(RegistroEstudio.objects.select_for_update().filter(
        pk__in=[caso.registro_estudio_id for caso in casos_bloqueados],
    ).order_by('registro_id', 'pk'))
    if len(casos_bloqueados) != len(caso_por_id):
        raise ValidationError('Un caso cambio mientras se preparaba la aplicacion; no se modifico ningun registro.')
    # Volver a validar después de adquirir los locks y justo antes de escribir.
    plan_actual = preparar_aplicacion_cantidad_doppler_mmii(caso_ids=list(caso_por_id), usuario=usuario)
    snapshots_plan = [
        (item['registro'].pk, str(item['monto_anterior']), str(item['monto_nuevo']), item['cantidades'])
        for item in plan
    ]
    snapshots_actuales = [
        (item['registro'].pk, str(item['monto_anterior']), str(item['monto_nuevo']), item['cantidades'])
        for item in plan_actual
    ]
    if snapshots_actuales != snapshots_plan:
        raise ValidationError('Los datos o importes cambiaron desde el preview; no se aplico ningun ajuste.')

    lote_id = uuid4()
    fecha_modificacion = timezone.now()
    resultados = []
    for item in plan_actual:
        registro = item['registro']
        monto_anterior = item['monto_anterior']
        cantidad_por_linea = item['cantidades']
        for caso in item['casos']:
            relacion_id = caso.registro_estudio_id
            relacion = RegistroEstudio.objects.get(pk=relacion_id)
            HistorialAjusteCantidadDopplerMMII.objects.create(
                registro=registro,
                registro_estudio=relacion,
                registro_estudio_id_origen=relacion_id,
                revision=caso,
                lote_id=lote_id,
                accion=HistorialAjusteCantidadDopplerMMII.ACCION_APLICAR,
                version_regla=caso.version_regla,
                cantidad_declarada=relacion.cantidad,
                cantidad_liquidable_anterior=relacion.cantidad_liquidable,
                cantidad_liquidable_nueva=cantidad_por_linea[relacion_id],
                monto_registro_anterior=monto_anterior,
                monto_registro_nuevo=item['monto_nuevo'],
                motivo=motivo,
                realizado_por=usuario,
            )
            RegistroEstudio.objects.filter(pk=relacion_id).update(
                cantidad_liquidable=cantidad_por_linea[relacion_id],
            )
        registro.monto_calculado = item['monto_nuevo']
        registro.modificado_por = usuario
        registro.fecha_modificacion = fecha_modificacion
        registro.motivo_modificacion = (
            f'Ajuste auditado de cantidad Doppler. Lote {lote_id}. '
            f'Monto ${monto_anterior} -> ${item["monto_nuevo"]}. {motivo}'
        )
        registro.save(update_fields=[
            'monto_calculado', 'modificado_por', 'fecha_modificacion', 'motivo_modificacion',
        ])
        resultados.append({
            'registro_id': registro.pk,
            'monto_anterior': str(monto_anterior),
            'monto_nuevo': str(item['monto_nuevo']),
            'casos': [caso.pk for caso in item['casos']],
        })
        for caso in item['casos']:
            caso.estimacion_json = {
                **caso.estimacion_json,
                'ajuste_cantidad_aplicado': True,
                'lote_ajuste_cantidad': str(lote_id),
                'cantidad_liquidable': cantidad_por_linea[caso.registro_estudio_id],
                'monto_registro_tras_aplicar': str(item['monto_nuevo']),
            }
            caso.save(update_fields=['estimacion_json'])
    return {'lote_id': str(lote_id), 'registros': resultados, 'casos_aplicados': len(caso_por_id)}


def preparar_reversion_lote_cantidad_doppler_mmii(*, lote_id, usuario):
    if not usuario.is_authenticated or not (usuario.is_superuser or usuario.rol == 'jefe_servicio'):
        raise ValidationError('No tienes permisos para revertir ajustes Doppler.')
    try:
        lote_id = UUID(str(lote_id))
    except (TypeError, ValueError, AttributeError):
        raise ValidationError('El identificador del lote no es valido.')
    eventos = list(HistorialAjusteCantidadDopplerMMII.objects.filter(
        lote_id=lote_id,
        accion=HistorialAjusteCantidadDopplerMMII.ACCION_APLICAR,
    ).select_related(
        'registro__medico', 'registro__sesion_contable', 'registro_estudio', 'revision',
    ).order_by('registro_id', 'registro_estudio_id_origen', 'pk'))
    if not eventos:
        raise ValidationError('No hay aplicaciones Doppler para ese lote.')
    errores = []
    eventos_por_registro = defaultdict(list)
    for evento in eventos:
        registro = evento.registro
        relacion = evento.registro_estudio
        if HistorialAjusteCantidadDopplerMMII.objects.filter(evento_origen=evento).exists():
            errores.append(f'La linea #{evento.registro_estudio_id_origen} ya fue revertida.')
        ultimo = HistorialAjusteCantidadDopplerMMII.objects.filter(
            registro_estudio_id_origen=evento.registro_estudio_id_origen,
        ).order_by('-fecha_evento', '-pk').first()
        if ultimo != evento:
            errores.append(f'La linea #{evento.registro_estudio_id_origen} tiene un evento posterior.')
        if (
            relacion.cantidad != evento.cantidad_declarada
            or relacion.cantidad_liquidable != evento.cantidad_liquidable_nueva
        ):
            errores.append(f'La cantidad actual de la linea #{evento.registro_estudio_id_origen} cambio.')
        if registro.monto_calculado != evento.monto_registro_nuevo:
            errores.append(f'El monto actual del registro #{registro.pk} cambio.')
        if registro.anulado or registro.medico.rol == 'medico_residente':
            errores.append(f'El registro #{registro.pk} ya no admite este ajuste economico.')
        if not registro.sesion_contable_id or registro.sesion_contable.estado not in {'ABIERTA', 'REVISION'}:
            errores.append(f'La sesion del registro #{registro.pk} no permite revertir.')
        eventos_por_registro[registro.pk].append(evento)
    if errores:
        raise ValidationError('No se puede revertir el lote: ' + ' '.join(errores))
    resultado = []
    for registro_id, eventos_registro in eventos_por_registro.items():
        registro = eventos_registro[0].registro
        montos_anteriores = {evento.monto_registro_anterior for evento in eventos_registro}
        montos_nuevos = {evento.monto_registro_nuevo for evento in eventos_registro}
        if len(montos_anteriores) != 1 or len(montos_nuevos) != 1:
            raise ValidationError(f'El historial del registro #{registro_id} no tiene snapshots coherentes.')
        resultado.append({
            'registro': registro,
            'eventos': eventos_registro,
            'monto_actual': registro.monto_calculado,
            'monto_restaurado': eventos_registro[0].monto_registro_anterior,
        })
    return resultado


@transaction.atomic
def revertir_lote_cantidad_doppler_mmii(*, lote_id, usuario, motivo):
    motivo = (motivo or '').strip()
    if not motivo or len(motivo) > 2000:
        raise ValidationError('Indica un fundamento de hasta 2000 caracteres.')
    plan = preparar_reversion_lote_cantidad_doppler_mmii(lote_id=lote_id, usuario=usuario)
    evento_ids = [evento.pk for item in plan for evento in item['eventos']]
    sesion_ids = sorted({item['registro'].sesion_contable_id for item in plan})
    list(SesionContable.objects.select_for_update().filter(
        pk__in=sesion_ids,
    ).order_by('pk'))
    registro_ids = sorted(item['registro'].pk for item in plan)
    list(RegistroEstudiosPorMedico.objects.select_for_update().filter(
        pk__in=registro_ids,
    ).order_by('pk'))
    list(HistorialAjusteCantidadDopplerMMII.objects.select_for_update().filter(
        pk__in=evento_ids,
    ).order_by('pk'))
    list(RegistroEstudio.objects.select_for_update().filter(
        pk__in=[evento.registro_estudio_id for item in plan for evento in item['eventos']],
    ).order_by('registro_id', 'pk'))
    plan_actual = preparar_reversion_lote_cantidad_doppler_mmii(lote_id=lote_id, usuario=usuario)
    estado_plan = [
        (item['registro'].pk, str(item['monto_actual']), str(item['monto_restaurado']), [e.pk for e in item['eventos']])
        for item in plan
    ]
    estado_actual = [
        (item['registro'].pk, str(item['monto_actual']), str(item['monto_restaurado']), [e.pk for e in item['eventos']])
        for item in plan_actual
    ]
    if estado_actual != estado_plan:
        raise ValidationError('El lote cambio desde el preview; no se revirtio ningun ajuste.')

    lote_reversion = uuid4()
    fecha_modificacion = timezone.now()
    resultados = []
    for item in plan_actual:
        registro = item['registro']
        for evento in item['eventos']:
            relacion = evento.registro_estudio
            HistorialAjusteCantidadDopplerMMII.objects.create(
                registro=registro,
                registro_estudio=relacion,
                registro_estudio_id_origen=evento.registro_estudio_id_origen,
                revision=evento.revision,
                evento_origen=evento,
                lote_id=lote_reversion,
                accion=HistorialAjusteCantidadDopplerMMII.ACCION_REVERTIR,
                version_regla=evento.version_regla,
                cantidad_declarada=evento.cantidad_declarada,
                cantidad_liquidable_anterior=relacion.cantidad_liquidable,
                cantidad_liquidable_nueva=evento.cantidad_liquidable_anterior,
                monto_registro_anterior=registro.monto_calculado,
                monto_registro_nuevo=item['monto_restaurado'],
                motivo=motivo,
                realizado_por=usuario,
            )
            RegistroEstudio.objects.filter(pk=relacion.pk).update(
                cantidad_liquidable=evento.cantidad_liquidable_anterior,
            )
            estimacion = dict(evento.revision.estimacion_json)
            estimacion.update({
                'ajuste_cantidad_aplicado': False,
                'lote_reversion_cantidad': str(lote_reversion),
                'monto_registro_tras_revertir': str(item['monto_restaurado']),
            })
            evento.revision.estimacion_json = estimacion
            evento.revision.save(update_fields=['estimacion_json'])
        registro.monto_calculado = item['monto_restaurado']
        registro.modificado_por = usuario
        registro.fecha_modificacion = fecha_modificacion
        registro.motivo_modificacion = (
            f'Reversion auditada del ajuste Doppler {lote_id}. '
            f'Monto ${item["monto_actual"]} -> ${item["monto_restaurado"]}. {motivo}'
        )
        registro.save(update_fields=[
            'monto_calculado', 'modificado_por', 'fecha_modificacion', 'motivo_modificacion',
        ])
        resultados.append({
            'registro_id': registro.pk,
            'monto_anterior': str(item['monto_actual']),
            'monto_restaurado': str(item['monto_restaurado']),
            'eventos_revertidos': [evento.pk for evento in item['eventos']],
        })
    return {
        'lote_id': str(lote_id),
        'lote_reversion': str(lote_reversion),
        'registros': resultados,
        'eventos_revertidos': len(evento_ids),
    }


@transaction.atomic
def confirmar_lote_auditoria_doppler_mmii(*, caso_ids, fecha_desde, fecha_hasta,
                                         medico_id, evidencias, observacion, usuario):
    if not usuario.is_authenticated or not (
        usuario.is_superuser or usuario.rol in {'administrativo', 'jefe_servicio'}
    ):
        raise ValidationError('No tienes permisos para confirmar casos Doppler.')
    ids = set(caso_ids)
    if not ids or len(ids) > 200:
        raise ValidationError('Selecciona entre 1 y 200 casos.')
    casos_qs = RevisionAuditoriaDopplerMMII.objects.filter(
        pk__in=ids, registro__fecha_del_informe__range=(fecha_desde, fecha_hasta),
        registro__anulado=False,
    )
    if medico_id:
        casos_qs = casos_qs.filter(registro__medico_id=medico_id)
    registros_ids = sorted(set(casos_qs.values_list('registro_id', flat=True)))
    list(RegistroEstudiosPorMedico.objects.select_for_update().filter(
        pk__in=registros_ids,
    ).order_by('pk'))
    casos = list(RevisionAuditoriaDopplerMMII.objects.select_for_update().filter(
        pk__in=list(casos_qs.values_list('pk', flat=True)),
    ).order_by('registro_id', 'pk'))
    if len(casos) != len(ids) or not all(caso_doppler_permite_lote(caso) for caso in casos):
        raise ValidationError('La seleccion contiene casos fuera del filtro, resueltos o no elegibles.')
    actual = auditar_cantidad_doppler_mmii(fecha_desde=fecha_desde, fecha_hasta=fecha_hasta)
    lineas_actuales = {
        estudio['registro_estudio_id']: estudio
        for resultado in actual['resultados'] for estudio in resultado['estudios']
    }
    for caso in casos:
        linea = lineas_actuales.get(caso.registro_estudio_id_origen)
        if not linea or linea['posible_duplicado'] or linea['cantidad_declarada'] != 2 or linea['cantidad_esperada'] != 1:
            raise ValidationError('Un caso cambio o es posible duplicado. Revisa la seleccion antes de confirmar.')
        registro = RegistroEstudiosPorMedico.objects.get(pk=caso.registro_id)
        fuente = caso.datos_originales_json
        if fuente.get('fuente_calculo') and fuente['fuente_calculo'] != fuente_calculo_doppler_mmii(registro):
            raise ValidationError('Los datos del caso cambiaron desde la deteccion; revisalo individualmente.')
    for caso in casos:
        resolver_caso_auditoria_doppler_mmii(
            caso_id=caso.pk, decision='CONFIRMADO', evidencias=evidencias,
            observacion=observacion, usuario=usuario,
        )
    return len(casos)


def adjuntar_comparacion_doppler_mmii(registros):
    """Presentacion compartida por pantallas y exportaciones, sin recalcular al leer."""
    registros = list(registros)
    casos_por_registro = defaultdict(list)
    for caso in RevisionAuditoriaDopplerMMII.objects.filter(
        registro_id__in=[registro.pk for registro in registros],
        version_regla=RevisionAuditoriaDopplerMMII.VERSION_REGLA_V1,
    ).select_related('revisado_por').order_by('fecha_deteccion', 'pk'):
        casos_por_registro[caso.registro_id].append(caso)
    ajustes_por_registro = defaultdict(list)
    for relacion in RegistroEstudio.objects.filter(
        registro_id__in=[registro.pk for registro in registros],
        cantidad_liquidable__isnull=False,
        estudio__tipo='DOP',
    ).select_related('estudio').order_by('registro_id', 'pk'):
        ajustes_por_registro[relacion.registro_id].append({
            'estudio': relacion.estudio.nombre,
            'cantidad_declarada': relacion.cantidad,
            'cantidad_liquidable': relacion.cantidad_liquidable,
        })
    for registro in registros:
        casos = casos_por_registro[registro.pk]
        registro.auditorias_doppler_mmii = casos
        registro.ajustes_cantidad_doppler_mmii = ajustes_por_registro[registro.pk]
        registro.comparacion_doppler = None
        if not casos:
            continue
        ultimo = max(casos, key=lambda caso: (caso.fecha_revision or caso.fecha_deteccion, caso.pk))
        estimacion = dict(ultimo.estimacion_json)
        fuente_original = casos[0].datos_originales_json.get('monto_registrado')
        comparacion = {
            'monto_original': Decimal(fuente_original) if fuente_original is not None else None,
            'monto_vigente': registro.monto_calculado,
            'monto_estimado': None,
            'diferencia_estimada': None,
            'motivo': estimacion.get('motivo', 'Pendiente de estimacion.'),
            'parcial': estimacion.get('parcial', True),
            'informativo_residente': estimacion.get('informativo_residente', registro.medico.rol == 'medico_residente'),
            'estado': ' / '.join(dict.fromkeys(caso.get_estado_display() for caso in casos)),
            'sin_ajustes_propuestos': all(caso.estado == 'DESCARTADO' for caso in casos),
        }
        ids_confirmados = sorted(caso.pk for caso in casos if caso.estado == 'CONFIRMADO')
        if (
            estimacion.get('monto_estimado') is not None
            and ids_confirmados == sorted(estimacion.get('casos_confirmados', []))
            and registro.monto_calculado == comparacion['monto_original']
            and not registro.anulado
            and estimacion.get('fuente_calculo') == fuente_calculo_doppler_mmii(registro)
        ):
            comparacion['monto_estimado'] = Decimal(estimacion['monto_estimado'])
            comparacion['diferencia_estimada'] = Decimal(estimacion['diferencia_estimada'])
        elif estimacion.get('monto_estimado') is not None:
            comparacion['motivo'] = 'Los datos o decisiones actuales cambiaron; requiere revisar la estimacion.'
        registro.comparacion_doppler = comparacion
    return registros


def resumir_comparacion_doppler_mmii(registros):
    registros = list(registros)
    comparaciones = [registro.comparacion_doppler for registro in registros if registro.comparacion_doppler]
    disponibles = [item for item in comparaciones if item['monto_estimado'] is not None]
    diferencias = [item['diferencia_estimada'] for item in disponibles if not item['informativo_residente']]
    posible_debito = sum((max(diferencia, Decimal('0.00')) for diferencia in diferencias), Decimal('0.00'))
    posible_credito = sum((max(-diferencia, Decimal('0.00')) for diferencia in diferencias), Decimal('0.00'))
    pendientes_economicos = sum(
        not item['informativo_residente'] and not item.get('sin_ajustes_propuestos')
        and (item['monto_estimado'] is None or item['parcial'])
        for item in comparaciones
    )
    return {
        'registros_auditados': len(comparaciones),
        'registros_estimados': len(disponibles),
        'pendientes_estimacion': len(comparaciones) - len(disponibles),
        'diferencia_estimada': sum((
            item['diferencia_estimada'] for item in disponibles if not item['informativo_residente']
        ), Decimal('0.00')),
        'diferencia_informativa_residentes': sum((
            item['diferencia_estimada'] for item in disponibles if item['informativo_residente']
        ), Decimal('0.00')),
        'posible_debito': posible_debito,
        'posible_credito': posible_credito,
        'monto_proyectado_practicas': sum((registro.monto_calculado for registro in registros), Decimal('0.00')) - posible_debito + posible_credito,
        'pendientes_economicos': pendientes_economicos,
    }


def valores_proyeccion_doppler_mmii(registro):
    comparacion = registro.comparacion_doppler
    if not comparacion or comparacion['informativo_residente'] or comparacion.get('sin_ajustes_propuestos'):
        return [0, 0, float(registro.monto_calculado)]
    diferencia = comparacion['diferencia_estimada']
    if diferencia is None:
        return ['Pendiente', 'Pendiente', 'Pendiente']
    return [
        float(max(diferencia, Decimal('0.00'))),
        float(max(-diferencia, Decimal('0.00'))),
        float(registro.monto_calculado - diferencia),
    ]


COLUMNAS_COMPARACION_DOPPLER_MMII = [
    'Monto original auditoria Doppler', 'Monto estimado auditoria Doppler',
    'Diferencia estimada Doppler (no aplicada)', 'Estado auditoria Doppler',
    'Alcance estimacion Doppler', 'Debito Doppler aplicado',
]


def valores_comparacion_doppler_mmii(registro):
    comparacion = registro.comparacion_doppler
    if not comparacion:
        return ['', '', '', 'Sin auditoria', '', 'No aplicado']
    return [
        float(comparacion['monto_original']) if comparacion['monto_original'] is not None else 'Pendiente',
        float(comparacion['monto_estimado']) if comparacion['monto_estimado'] is not None else 'Pendiente',
        float(comparacion['diferencia_estimada']) if comparacion['diferencia_estimada'] is not None else 'Pendiente',
        comparacion['estado'],
        'Informativo sin debito (residente)' if comparacion['informativo_residente'] else (
            'Parcial: solo cantidades confirmadas' if comparacion['parcial'] else 'Cantidades confirmadas'
        ),
        'No aplicado',
    ]


def agregar_hoja_auditoria_doppler_mmii(wb, registros):
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    sheet = wb.create_sheet('Auditoria Doppler')
    sheet.append([
        'Registro', 'Caso', 'Profesional', 'Fecha informe', 'Practica',
        'Cantidad original', 'Cantidad segun regla', 'Estado',
        'Orden verificada', 'VisualMedical verificado', 'EGES consultado', 'NetTerm consultado',
        'Fundamento', 'Revisor', 'Fecha revision',
        'Monto original registro', 'Monto vigente registro', 'Monto estimado registro',
        'Diferencia estimada registro', 'Alcance', 'Motivo estimacion',
    ])
    for registro in registros:
        comparacion = registro.comparacion_doppler
        for indice, caso in enumerate(registro.auditorias_doppler_mmii):
            fuente = caso.datos_originales_json
            mostrar_importe = indice == 0
            sheet.append([
                registro.pk, caso.pk, registro.medico.get_full_name() or registro.medico.username,
                fuente.get('fecha_informe'), fuente.get('estudio'), fuente.get('cantidad_declarada'),
                fuente.get('cantidad_esperada'), caso.get_estado_display(),
                caso.orden_medica_verificada, caso.visualmedical_verificado,
                caso.eges_verificado, caso.netterm_verificado, caso.observacion,
                caso.revisado_por.get_full_name() or caso.revisado_por.username if caso.revisado_por else '',
                caso.fecha_revision.isoformat() if caso.fecha_revision else '',
                float(comparacion['monto_original']) if mostrar_importe and comparacion['monto_original'] is not None else '',
                float(registro.monto_calculado) if mostrar_importe else '',
                valores_comparacion_doppler_mmii(registro)[1] if mostrar_importe else '',
                valores_comparacion_doppler_mmii(registro)[2] if mostrar_importe else '',
                valores_comparacion_doppler_mmii(registro)[4], comparacion['motivo'],
            ])
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if 16 <= cell.column <= 19:
                cell.number_format = '$#,##0.00'
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for column in sheet.columns:
        sheet.column_dimensions[get_column_letter(column[0].column)].width = min(
            max(len(str(cell.value or '')) for cell in column) + 2, 45,
        )
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions


def _es_monto_cero(monto):
    try:
        return Decimal(str(monto or 0)) <= Decimal('0')
    except Exception:
        return True


def _debe_bloquear(tipo, estado_destino):
    # ABIERTA -> REVISION: solo advertencias
    if estado_destino == 'REVISION':
        return False

    # REVISION -> CERRADA: bloqueos estructurales
    if estado_destino == 'CERRADA':
        return tipo in {
            TIPO_SIN_REGISTRO_ESTUDIO,
            TIPO_MONTO_CERO_CON_ESTUDIOS,
            TIPO_GUARDIA_MONTO_INVALIDO,
        }

    # CERRADA -> FACTURADA: bloqueos financieros + sesión vacía
    if estado_destino == 'FACTURADA':
        return tipo in {
            TIPO_SESION_VACIA,
            TIPO_SIN_PRECIO_RESOLUBLE,
            TIPO_SIN_TARIFA_VIGENTE_GRUPO,
            TIPO_CONTEXTUAL_SIN_GRUPO,
            TIPO_CONTEXTUAL_SIN_TARIFA,
        }

    # FACTURADA -> PAGADA: sesión vacía o remanentes críticos
    if estado_destino == 'PAGADA':
        return tipo in {
            TIPO_SESION_VACIA,
            TIPO_SIN_REGISTRO_ESTUDIO,
            TIPO_MONTO_CERO_CON_ESTUDIOS,
            TIPO_GUARDIA_MONTO_INVALIDO,
            TIPO_SIN_PRECIO_RESOLUBLE,
            TIPO_SIN_TARIFA_VIGENTE_GRUPO,
            TIPO_CONTEXTUAL_SIN_GRUPO,
            TIPO_CONTEXTUAL_SIN_TARIFA,
        }

    return False


def _serializar_metadata_issue(metadata):
    serializada = {}
    for key, value in metadata.items():
        if hasattr(value, 'isoformat'):
            serializada[key] = value.isoformat()
        else:
            serializada[key] = value
    return serializada


def _agregar_issue(resultado, tipo, mensaje, estado_destino, **metadata):
    es_bloqueante = _debe_bloquear(tipo, estado_destino)
    destino = resultado['bloqueantes'] if es_bloqueante else resultado['advertencias']
    destino.append(mensaje)
    resultado.setdefault('items', []).append({
        'tipo': tipo,
        'mensaje': mensaje,
        'estado': 'bloqueante' if es_bloqueante else 'advertencia',
        **_serializar_metadata_issue(metadata),
    })


def evaluar_gate_consistencia_sesion(sesion, estado_destino):
    """
    Evalua consistencia administrativa de una sesion antes de transicionar de estado.

    Retorna:
      {
        'bloqueantes': [...],
        'advertencias': [...],
      }
    """
    resultado = {'bloqueantes': [], 'advertencias': [], 'items': []}

    practicas = list(
        sesion.practicas.filter(anulado=False).select_related('medico', 'sesion_contable').prefetch_related('registroestudio_set__estudio')
    )
    guardias = list(sesion.guardias_pasivas.all())

    if estado_destino in {'FACTURADA', 'PAGADA'} and not practicas and not guardias:
        _agregar_issue(
            resultado,
            TIPO_SESION_VACIA,
            'Sesion vacia: no hay practicas ni guardias para transicionar.',
            estado_destino,
        )

    for registro in practicas:
        relaciones = list(registro.registroestudio_set.all())

        if not relaciones:
            _agregar_issue(
                resultado,
                TIPO_SIN_REGISTRO_ESTUDIO,
                f'Registro #{registro.pk} sin estudios asociados en RegistroEstudio.',
                estado_destino,
                registro_id=registro.pk,
                fecha=registro.fecha_del_informe,
            )
            continue

        if _es_monto_cero(registro.monto_calculado):
            _agregar_issue(
                resultado,
                TIPO_MONTO_CERO_CON_ESTUDIOS,
                f'Registro #{registro.pk} con estudios asociados y monto_calculado <= 0.',
                estado_destino,
                registro_id=registro.pk,
                fecha=registro.fecha_del_informe,
            )

        for rel in relaciones:
            estudio = rel.estudio
            fecha_ref = registro.fecha_del_informe
            contexto = (rel.contexto or 'SERVICIO').upper()

            if estudio.grupo_tarifario_id:
                tarifa_base = estudio.grupo_tarifario.get_tarifa_vigente(fecha=fecha_ref)
                if not tarifa_base:
                    _agregar_issue(
                        resultado,
                        TIPO_SIN_TARIFA_VIGENTE_GRUPO,
                        (
                            f'Registro #{registro.pk} estudio {estudio.nombre}: '
                            f'grupo {estudio.grupo_tarifario.codigo} sin tarifa vigente para {fecha_ref}.'
                        ),
                        estado_destino,
                        registro_id=registro.pk,
                        estudio_id=estudio.pk,
                        grupo_id=estudio.grupo_tarifario_id,
                        grupo_codigo=estudio.grupo_tarifario.codigo,
                        fecha=fecha_ref,
                    )

                if contexto in {'LECHO', 'QUIROFANO'}:
                    codigo_ctx = f'{estudio.grupo_tarifario.codigo}_{contexto}'
                    grupo_ctx = GrupoTarifario.objects.filter(codigo=codigo_ctx, activo=True).first()
                    if not grupo_ctx:
                        _agregar_issue(
                            resultado,
                            TIPO_CONTEXTUAL_SIN_GRUPO,
                            (
                                f'Registro #{registro.pk} estudio {estudio.nombre}: '
                                f'contexto {contexto} sin grupo contextual {codigo_ctx}.'
                            ),
                            estado_destino,
                            registro_id=registro.pk,
                            estudio_id=estudio.pk,
                            grupo_codigo=codigo_ctx,
                            contexto=contexto,
                            fecha=fecha_ref,
                        )
                    else:
                        tarifa_ctx = grupo_ctx.get_tarifa_vigente(fecha=fecha_ref)
                        if not tarifa_ctx:
                            _agregar_issue(
                                resultado,
                                TIPO_CONTEXTUAL_SIN_TARIFA,
                                (
                                    f'Registro #{registro.pk} estudio {estudio.nombre}: '
                                    f'grupo contextual {codigo_ctx} sin tarifa vigente para {fecha_ref}.'
                                ),
                                estado_destino,
                                registro_id=registro.pk,
                                estudio_id=estudio.pk,
                                grupo_id=grupo_ctx.pk,
                                grupo_codigo=codigo_ctx,
                                contexto=contexto,
                                fecha=fecha_ref,
                            )
            else:
                precio_legado = estudio.precio_para_os(
                    registro.tipo_obra_social,
                    fecha=fecha_ref,
                    contexto=contexto,
                )
                if _es_monto_cero(precio_legado):
                    _agregar_issue(
                        resultado,
                        TIPO_SIN_PRECIO_RESOLUBLE,
                        (
                            f'Registro #{registro.pk} estudio {estudio.nombre}: '
                            'sin grupo y sin precio legado valido (>0).'
                        ),
                        estado_destino,
                        registro_id=registro.pk,
                        estudio_id=estudio.pk,
                        fecha=fecha_ref,
                    )
                else:
                    _agregar_issue(
                        resultado,
                        TIPO_SIN_GRUPO_CON_FALLBACK,
                        (
                            f'Registro #{registro.pk} estudio {estudio.nombre}: '
                            'sin grupo tarifario (se usa fallback legado).'
                        ),
                        estado_destino,
                        registro_id=registro.pk,
                        estudio_id=estudio.pk,
                        fecha=fecha_ref,
                    )

    for guardia in guardias:
        if _es_monto_cero(guardia.monto):
            _agregar_issue(
                resultado,
                TIPO_GUARDIA_MONTO_INVALIDO,
                f'Guardia #{guardia.pk} con monto <= 0.',
                estado_destino,
                guardia_id=guardia.pk,
                fecha=getattr(guardia, 'fecha_guardia', None),
            )

    return resultado


def _nivel_alerta(valor, umbral_amarillo, umbral_rojo):
    if valor >= umbral_rojo:
        return 'roja'
    if valor >= umbral_amarillo:
        return 'amarilla'
    return None


def _agregar_alerta(alertas, tipo, valor, umbral_amarillo, umbral_rojo, mensaje):
    nivel = _nivel_alerta(valor, umbral_amarillo, umbral_rojo)
    if not nivel:
        return

    alertas.append({
        'tipo': tipo,
        'severidad': nivel,
        'valor': valor,
        'umbral': umbral_rojo if nivel == 'roja' else umbral_amarillo,
        'mensaje': mensaje,
    })


def auditar_residentes_eco_por_sesion(sesion):
    """
    Auditoría administrativa de patrones ECO en residentes para una sesión contable.

    Retorno (contrato):
      {
        'sesion_id': int,
        'periodo': str,
        'total_residentes': int,
        'residentes_con_alertas': int,
        'alertas_rojas': int,
        'alertas_amarillas': int,
        'items': [ ... ],
        'top_alertas': [ ... ],
      }
    """
    meses = ['', 'Enero', 'Febrero', 'Marzo', 'Abril', 'Mayo', 'Junio',
             'Julio', 'Agosto', 'Septiembre', 'Octubre', 'Noviembre', 'Diciembre']
    periodo = f"{meses[sesion.mes] if 1 <= sesion.mes <= 12 else sesion.mes} {sesion.año}"

    resultado = {
        'sesion_id': sesion.id,
        'periodo': periodo,
        'total_residentes': 0,
        'residentes_con_alertas': 0,
        'alertas_rojas': 0,
        'alertas_amarillas': 0,
        'items': [],
        'top_alertas': [],
    }

    practicas = list(
        sesion.practicas.filter(anulado=False).select_related('medico').prefetch_related('registroestudio_set__estudio')
        .filter(medico__rol__in=ROLES_AUDITORIA_RESIDENCIA_ECO)
    )
    if not practicas:
        return resultado

    registro_ids = [registro.pk for registro in practicas]
    revisiones_eges_requieren_correccion = set()
    revisiones_eges = (
        RevisionCruceEgesRegistro.objects
        .filter(sesion_contable=sesion, registro_id__in=registro_ids)
        .order_by('registro_id', '-fecha_revision')
    )
    ultimas_revisiones_eges = {}
    for revision in revisiones_eges:
        ultimas_revisiones_eges.setdefault(revision.registro_id, revision)
    revisiones_eges_requieren_correccion = {
        registro_id
        for registro_id, revision in ultimas_revisiones_eges.items()
        if revision.estado == RevisionCruceEgesRegistro.ESTADO_REQUIERE_CORRECCION
    }

    fechas_locales = []
    for registro in practicas:
        dt = registro.fecha_registro
        if not dt:
            continue
        if timezone.is_naive(dt):
            dt = timezone.make_aware(dt, timezone.get_current_timezone())
        fechas_locales.append(timezone.localtime(dt).date())

    feriados = set(Feriado.objects.filter(fecha__in=fechas_locales).values_list('fecha', flat=True))

    acumulado = {}

    for registro in practicas:
        relaciones = list(registro.registroestudio_set.all())
        if not relaciones:
            continue

        if not any(rel.estudio.tipo in TIPOS_AUDITORIA_RESIDENCIA_ECO for rel in relaciones):
            continue

        medico = registro.medico
        item = acumulado.setdefault(
            medico.id,
            {
                'medico_id': medico.id,
                'medico_nombre': medico.get_full_name() or medico.username,
                'rol': medico.rol,
                'total_eco': 0,
                'intra': 0,
                'extra': 0,
                'proporcion_extra': 0,
                'nocturnos': 0,
                'finde_feriado': 0,
                'post_17': 0,
                'max_eco_dia': 0,
                'dias_pico': [],
                'severidad': 'ok',
                'alertas': [],
                '_eco_por_dia': defaultdict(int),
                '_registros_eco': [],
            },
        )

        item['total_eco'] += 1
        if registro.horario == 'INTRA':
            item['intra'] += 1
        elif registro.horario == 'EXTRA':
            item['extra'] += 1

        dt = registro.fecha_registro
        if not dt:
            continue
        if timezone.is_naive(dt):
            dt = timezone.make_aware(dt, timezone.get_current_timezone())
        dt = timezone.localtime(dt)

        fecha_local = dt.date()
        hora_local = dt.hour
        weekday = fecha_local.weekday()
        es_finde = weekday >= 5
        es_feriado = fecha_local in feriados

        if hora_local >= AUDIT_HORA_NOCTURNA_DESDE or hora_local < AUDIT_HORA_NOCTURNA_HASTA:
            item['nocturnos'] += 1

        if es_finde or es_feriado:
            item['finde_feriado'] += 1
        elif hora_local >= AUDIT_HORA_POST_17:
                item['post_17'] += 1

        item['_eco_por_dia'][fecha_local] += 1
        motivos_registro = []
        if registro.pk in revisiones_eges_requieren_correccion:
            motivos_registro.append('EGES requiere correccion')
        if registro.horario == 'EXTRA':
            motivos_registro.append('EXTRA')
        if hora_local >= AUDIT_HORA_NOCTURNA_DESDE or hora_local < AUDIT_HORA_NOCTURNA_HASTA:
            motivos_registro.append('Nocturno')
        if es_finde or es_feriado:
            motivos_registro.append('Finde/Feriado')
        elif hora_local >= AUDIT_HORA_POST_17:
            motivos_registro.append('Post 17')

        item['_registros_eco'].append({
            'registro_id': registro.pk,
            'fecha_informe': registro.fecha_del_informe.isoformat() if registro.fecha_del_informe else '',
            'fecha_carga': dt.isoformat(),
            'fecha_local': fecha_local.isoformat(),
            'estudios': ', '.join(
                rel.estudio.nombre
                for rel in relaciones
                if rel.estudio.tipo in TIPOS_AUDITORIA_RESIDENCIA_ECO
            ),
            'paciente': f'{registro.apellido_paciente}, {registro.nombre_paciente}',
            'dni_paciente': registro.dni_paciente,
            'horario': registro.horario,
            'monto_calculado': registro.monto_calculado,
            'motivos': motivos_registro,
        })

    items = []
    for item in acumulado.values():
        total_eco = item['total_eco']
        if total_eco <= 0:
            continue

        item['proporcion_extra'] = item['extra'] / total_eco

        eco_por_dia = item.pop('_eco_por_dia')
        registros_eco = item.pop('_registros_eco')
        item['max_eco_dia'] = max(eco_por_dia.values()) if eco_por_dia else 0
        item['dias_pico'] = sorted(
            [
                fecha.isoformat()
                for fecha, cantidad in eco_por_dia.items()
                if cantidad >= AUDIT_MAX_ECO_DIA_AMARILLA
            ]
        )

        _agregar_alerta(
            item['alertas'],
            'extra_mensual',
            item['extra'],
            AUDIT_EXTRA_MENSUAL_AMARILLA,
            AUDIT_EXTRA_MENSUAL_ROJA,
            'Cantidad de registros EXTRA mensual elevada',
        )
        _agregar_alerta(
            item['alertas'],
            'proporcion_extra',
            item['proporcion_extra'],
            AUDIT_PROP_EXTRA_AMARILLA,
            AUDIT_PROP_EXTRA_ROJA,
            'Proporción EXTRA elevada',
        )
        _agregar_alerta(
            item['alertas'],
            'nocturnos',
            item['nocturnos'],
            AUDIT_NOCTURNOS_AMARILLA,
            AUDIT_NOCTURNOS_ROJA,
            'Cantidad de registros nocturnos elevada',
        )
        _agregar_alerta(
            item['alertas'],
            'finde_feriado',
            item['finde_feriado'],
            AUDIT_FINDE_FERIADO_AMARILLA,
            AUDIT_FINDE_FERIADO_ROJA,
            'Cantidad de registros en sábado/domingo/feriado elevada',
        )
        _agregar_alerta(
            item['alertas'],
            'volumen_diario',
            item['max_eco_dia'],
            AUDIT_MAX_ECO_DIA_AMARILLA,
            AUDIT_MAX_ECO_DIA_ROJA,
            'Volumen diario ECO inusual',
        )
        registros_eges_requieren_correccion = sum(
            1
            for registro_alerta in registros_eco
            if 'EGES requiere correccion' in registro_alerta.get('motivos', [])
        )
        _agregar_alerta(
            item['alertas'],
            'eges_requiere_correccion',
            registros_eges_requieren_correccion,
            1,
            1,
            'Cruce EGES requiere correccion',
        )

        if any(alerta['severidad'] == 'roja' for alerta in item['alertas']):
            item['severidad'] = 'roja'
        elif item['alertas']:
            item['severidad'] = 'amarilla'
        else:
            item['severidad'] = 'ok'

        dias_pico = set(item['dias_pico'])
        for registro_alerta in registros_eco:
            if registro_alerta['fecha_local'] in dias_pico and 'Dia pico' not in registro_alerta['motivos']:
                registro_alerta['motivos'].append('Dia pico')

        tipos_alerta = {alerta['tipo'] for alerta in item['alertas']}
        registros_alerta = []
        for registro_alerta in registros_eco:
            motivos = set(registro_alerta['motivos'])
            aporta = (
                ({'extra_mensual', 'proporcion_extra'} & tipos_alerta and 'EXTRA' in motivos)
                or ('nocturnos' in tipos_alerta and 'Nocturno' in motivos)
                or ('finde_feriado' in tipos_alerta and 'Finde/Feriado' in motivos)
                or ('volumen_diario' in tipos_alerta and 'Dia pico' in motivos)
                or ('eges_requiere_correccion' in tipos_alerta and 'EGES requiere correccion' in motivos)
            )
            if aporta:
                registros_alerta.append(registro_alerta)

        registros_alerta.sort(key=lambda registro: registro['fecha_carga'], reverse=True)
        item['registros_alerta'] = registros_alerta
        item['registros_alerta_total'] = len(registros_alerta)

        items.append(item)

    orden_severidad = {'roja': 0, 'amarilla': 1, 'ok': 2}
    items.sort(key=lambda x: (orden_severidad.get(x['severidad'], 9), x['medico_nombre']))

    resultado['items'] = items
    resultado['total_residentes'] = len(items)
    resultado['residentes_con_alertas'] = len([i for i in items if i['severidad'] in {'roja', 'amarilla'}])
    resultado['alertas_rojas'] = sum(1 for i in items for a in i['alertas'] if a['severidad'] == 'roja')
    resultado['alertas_amarillas'] = sum(1 for i in items for a in i['alertas'] if a['severidad'] == 'amarilla')

    top = [i for i in items if i['severidad'] in {'roja', 'amarilla'}]
    top.sort(key=lambda x: (orden_severidad[x['severidad']], -len(x['alertas']), x['medico_nombre']))
    resultado['top_alertas'] = [
        {
            'medico_id': i['medico_id'],
            'medico_nombre': i['medico_nombre'],
            'severidad': i['severidad'],
            'cantidad_alertas': len(i['alertas']),
        }
        for i in top[:3]
    ]

    return resultado


def resumir_pendientes_auditoria_eco(auditoria):
    """
    Agrega una lectura operativa de pendientes sin alterar la auditoria bruta.

    La deteccion de patrones ECO conserva todos los hallazgos. Para paneles de
    cierre, en cambio, solo debe quedar pendiente lo que todavia requiere accion:
    registros sin revision o con revision "requiere correccion" sin ajuste PACS.
    """
    from .models import (
        CorreccionPacsRegistro,
        RevisionAuditoriaEcoRegistro,
        RevisionCruceEgesRegistro,
    )

    registro_ids = [
        registro_alerta.get('registro_id')
        for item in auditoria.get('items', [])
        for registro_alerta in item.get('registros_alerta', [])
        if registro_alerta.get('registro_id')
    ]
    registro_ids = list(dict.fromkeys(registro_ids))

    revisiones_por_registro = {}
    if registro_ids:
        revisiones = (
            RevisionAuditoriaEcoRegistro.objects
            .filter(registro_id__in=registro_ids)
            .order_by('registro_id', '-fecha_revision')
        )
        for revision in revisiones:
            revisiones_por_registro.setdefault(revision.registro_id, revision)

    revisiones_eges_por_registro = {}
    if registro_ids:
        revisiones_eges = (
            RevisionCruceEgesRegistro.objects
            .filter(registro_id__in=registro_ids)
            .order_by('registro_id', '-fecha_revision')
        )
        for revision in revisiones_eges:
            revisiones_eges_por_registro.setdefault(revision.registro_id, revision)

    registros_con_correccion = set()
    if registro_ids:
        registros_con_correccion = set(
            CorreccionPacsRegistro.objects
            .filter(registro_id__in=registro_ids)
            .values_list('registro_id', flat=True)
        )

    items_pendientes = []
    alertas_rojas_pendientes = 0
    alertas_amarillas_pendientes = 0
    registros_pendientes_total = 0

    for item in auditoria.get('items', []):
        registros_pendientes = []
        for registro_alerta in item.get('registros_alerta', []):
            registro_id = registro_alerta.get('registro_id')
            revision = revisiones_por_registro.get(registro_id)
            revision_eges = revisiones_eges_por_registro.get(registro_id)
            tiene_correccion = registro_id in registros_con_correccion

            pendiente = (
                revision is None
                or (
                    revision.estado == RevisionAuditoriaEcoRegistro.ESTADO_REQUIERE_CORRECCION
                    and not tiene_correccion
                )
            )
            if revision is None and revision_eges:
                pendiente = revision_eges.estado == RevisionCruceEgesRegistro.ESTADO_REQUIERE_CORRECCION
            registro_alerta['auditoria_eco_pendiente'] = pendiente
            if pendiente:
                registros_pendientes.append(registro_alerta)

        item['registros_alerta_pendientes'] = registros_pendientes
        item['registros_alerta_pendientes_total'] = len(registros_pendientes)
        item['auditoria_eco_pendiente'] = bool(registros_pendientes)

        if registros_pendientes:
            items_pendientes.append(item)
            registros_pendientes_total += len(registros_pendientes)
            alertas_rojas_pendientes += sum(
                1 for alerta in item.get('alertas', [])
                if alerta.get('severidad') == 'roja'
            )
            alertas_amarillas_pendientes += sum(
                1 for alerta in item.get('alertas', [])
                if alerta.get('severidad') == 'amarilla'
            )

    orden_severidad = {'roja': 0, 'amarilla': 1, 'ok': 2}
    top_pendientes = sorted(
        items_pendientes,
        key=lambda item: (
            orden_severidad.get(item.get('severidad'), 9),
            -item.get('registros_alerta_pendientes_total', 0),
            item.get('medico_nombre') or '',
        ),
    )

    auditoria['items_pendientes'] = items_pendientes
    auditoria['residentes_con_alertas_pendientes'] = len(items_pendientes)
    auditoria['registros_alerta_pendientes_total'] = registros_pendientes_total
    auditoria['alertas_rojas_pendientes'] = alertas_rojas_pendientes
    auditoria['alertas_amarillas_pendientes'] = alertas_amarillas_pendientes
    auditoria['top_alertas_pendientes'] = [
        {
            'medico_id': item.get('medico_id'),
            'medico_nombre': item.get('medico_nombre'),
            'severidad': item.get('severidad'),
            'cantidad_alertas': len(item.get('alertas', [])),
            'registros_pendientes': item.get('registros_alerta_pendientes_total', 0),
        }
        for item in top_pendientes[:3]
    ]

    return auditoria
