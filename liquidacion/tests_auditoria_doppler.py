from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.test import TestCase
from django.test import override_settings
from django.urls import reverse
from django.utils import timezone

from .models import (
    Estudios,
    GrupoTarifario,
    GuardiaPasiva,
    HistorialRevisionAuditoriaDopplerMMII,
    RegistroEstudio,
    RegistroEstudiosPorMedico,
    RevisionAuditoriaDopplerMMII,
    SesionContable,
    SolicitudRevisionHorarioRegistro,
)
from .services_auditoria import (
    adjuntar_comparacion_doppler_mmii,
    auditar_cantidad_doppler_mmii,
    confirmar_lote_auditoria_doppler_mmii,
    crear_casos_auditoria_doppler_mmii,
    resolver_caso_auditoria_doppler_mmii,
    resumir_comparacion_doppler_mmii,
    valores_proyeccion_doppler_mmii,
)


User = get_user_model()


class AuditoriaCantidadDopplerMMIITest(TestCase):
    def setUp(self):
        self.medico = User.objects.create_user(
            username='jefe_doppler_auditoria',
            first_name='Angel',
            last_name='Gavilanes Ibarra',
            rol='jefe_residentes',
            perfil_completo=True,
        )
        self.sesion = SesionContable.objects.create(
            mes=6,
            año=2026,
            estado='CERRADA',
        )
        self.grupo = GrupoTarifario.objects.create(
            codigo='DOP_MMII_AUDIT',
            nombre='Doppler MMII auditoria',
            modalidad='DOP',
            activo=True,
        )
        self.arterial = Estudios.objects.create(
            codigo='900046/0',
            nombre='Ecodoppler arterial MM inferiores',
            tipo='DOP',
            grupo_tarifario=self.grupo,
            conteo_regiones=1,
            conteo_regiones_default=1,
            precio_cober=Decimal('24200.00'),
            precio_otras_os=Decimal('24200.00'),
            activo=True,
        )
        self.venoso = Estudios.objects.create(
            codigo='DOP-VEN-MMII',
            nombre='Doppler venoso de miembros inferiores',
            tipo='DOP',
            grupo_tarifario=self.grupo,
            conteo_regiones=1,
            conteo_regiones_default=1,
            precio_cober=Decimal('24200.00'),
            precio_otras_os=Decimal('24200.00'),
            activo=True,
        )

    def _crear_registro(self, *, rol='jefe_residentes'):
        if rol != self.medico.rol:
            self.medico.rol = rol
            self.medico.save(update_fields=['rol'])
        return RegistroEstudiosPorMedico.objects.create(
            sesion_contable=self.sesion,
            medico=self.medico,
            nombre_paciente='Claudia Vanesa',
            apellido_paciente='Sanchez Barrios',
            dni_paciente='95951354',
            fecha_del_informe=date(2026, 6, 10),
            tipo_obra_social='COBER',
            horario='EXTRA',
            cantidad_regiones=4,
            monto_calculado=Decimal('48400.00'),
        )

    def test_normaliza_arterial_y_venoso_sin_modificar_datos_originales(self):
        registro = self._crear_registro()
        relacion_arterial = RegistroEstudio.objects.create(
            registro=registro,
            estudio=self.arterial,
            cantidad=2,
            contexto='SERVICIO',
        )
        relacion_venoso = RegistroEstudio.objects.create(
            registro=registro,
            estudio=self.venoso,
            cantidad=2,
            contexto='SERVICIO',
        )
        RegistroEstudiosPorMedico.objects.filter(pk=registro.pk).update(
            cantidad_regiones=4,
            monto_calculado=Decimal('48400.00'),
        )

        resultado = auditar_cantidad_doppler_mmii(
            fecha_desde=date(2026, 6, 1),
            fecha_hasta=date(2026, 9, 30),
        )

        fila = resultado['resultados'][0]
        self.assertEqual(fila['cantidad_regiones_declarada'], 4)
        self.assertEqual(fila['cantidad_regiones_esperada'], 2)
        self.assertEqual(fila['monto_registrado'], Decimal('48400.00'))
        self.assertEqual(fila['diferencia_cantidad'], 2)
        self.assertEqual(
            [item['cantidad_esperada'] for item in fila['estudios']],
            [1, 1],
        )
        self.assertEqual(
            [item['registro_estudio_id'] for item in fila['estudios']],
            [relacion_arterial.pk, relacion_venoso.pk],
        )
        self.assertEqual(fila['tratamiento_economico'], 'IMPACTO_POTENCIAL_NO_APLICADO')

        relacion_arterial.refresh_from_db()
        relacion_venoso.refresh_from_db()
        registro.refresh_from_db()
        self.assertEqual(relacion_arterial.cantidad, 2)
        self.assertEqual(relacion_venoso.cantidad, 2)
        self.assertEqual(registro.cantidad_regiones, 4)
        self.assertEqual(registro.monto_calculado, Decimal('48400.00'))

    def test_residente_solo_recibe_informe_sin_propuesta_de_debito(self):
        registro = self._crear_registro(rol='medico_residente')
        RegistroEstudio.objects.create(
            registro=registro,
            estudio=self.arterial,
            cantidad=2,
            contexto='SERVICIO',
        )
        RegistroEstudiosPorMedico.objects.filter(pk=registro.pk).update(cantidad_regiones=2)

        resultado = auditar_cantidad_doppler_mmii(
            fecha_desde=date(2026, 6, 1),
            fecha_hasta=date(2026, 9, 30),
        )

        self.assertEqual(
            resultado['resultados'][0]['tratamiento_economico'],
            'INFORMATIVO_SIN_DEBITO',
        )

    def test_no_clasifica_doppler_sin_referencia_a_mmii(self):
        registro = self._crear_registro()
        estudio_carotideo = Estudios.objects.create(
            codigo='DOP-CAROTIDEO',
            nombre='Doppler arterial carotideo',
            tipo='DOP',
            conteo_regiones=1,
            conteo_regiones_default=1,
            precio_cober=Decimal('1000.00'),
            precio_otras_os=Decimal('1000.00'),
            activo=True,
        )
        RegistroEstudio.objects.create(
            registro=registro,
            estudio=estudio_carotideo,
            cantidad=2,
        )

        resultado = auditar_cantidad_doppler_mmii(
            fecha_desde=date(2026, 6, 1),
            fecha_hasta=date(2026, 9, 30),
        )

        self.assertEqual(resultado['total_registros'], 0)

    def test_un_estudio_combinado_arterial_y_venoso_equivale_a_dos_practicas(self):
        registro = self._crear_registro()
        estudio_combinado = Estudios.objects.create(
            codigo='DOP-ART-VEN-MMII',
            nombre='Ecodoppler arterial y venoso MMII',
            tipo='DOP',
            grupo_tarifario=self.grupo,
            conteo_regiones=1,
            conteo_regiones_default=1,
            precio_cober=Decimal('48400.00'),
            precio_otras_os=Decimal('48400.00'),
            activo=True,
        )
        RegistroEstudio.objects.create(
            registro=registro,
            estudio=estudio_combinado,
            cantidad=1,
        )

        resultado = auditar_cantidad_doppler_mmii(
            fecha_desde=date(2026, 6, 1),
            fecha_hasta=date(2026, 9, 30),
        )

        estudio_auditado = resultado['resultados'][0]['estudios'][0]
        self.assertEqual(estudio_auditado['clasificacion'], 'ARTERIAL_Y_VENOSO_MMII')
        self.assertEqual(estudio_auditado['cantidad_esperada'], 2)

    def test_marca_posible_duplicado_sin_fusionar_registros(self):
        primer_registro = self._crear_registro()
        RegistroEstudio.objects.create(
            registro=primer_registro,
            estudio=self.arterial,
            cantidad=1,
        )
        segundo_registro = RegistroEstudiosPorMedico.objects.create(
            sesion_contable=self.sesion,
            medico=self.medico,
            nombre_paciente='Claudia Vanesa',
            apellido_paciente='Sanchez Barrios',
            dni_paciente='95.951.354',
            fecha_del_informe=primer_registro.fecha_del_informe,
            tipo_obra_social='COBER',
            horario='EXTRA',
            cantidad_regiones=1,
            monto_calculado=Decimal('24200.00'),
        )
        RegistroEstudio.objects.create(
            registro=segundo_registro,
            estudio=self.arterial,
            cantidad=1,
        )

        resultado = auditar_cantidad_doppler_mmii(
            fecha_desde=date(2026, 6, 1),
            fecha_hasta=date(2026, 9, 30),
        )

        self.assertEqual(resultado['grupos_posible_duplicado'], 1)
        self.assertEqual(len(resultado['resultados']), 2)
        for registro in resultado['resultados']:
            estudio = registro['estudios'][0]
            self.assertTrue(estudio['posible_duplicado'])
            self.assertEqual(estudio['cantidad_en_grupo_duplicado'], 2)
        self.assertEqual(RegistroEstudiosPorMedico.objects.count(), 2)
        self.assertEqual(RegistroEstudio.objects.count(), 2)

    def test_no_marca_art_and_venous_as_duplicate_together(self):
        registro_arterial = self._crear_registro()
        RegistroEstudio.objects.create(
            registro=registro_arterial,
            estudio=self.arterial,
            cantidad=1,
        )
        registro_venoso = RegistroEstudiosPorMedico.objects.create(
            sesion_contable=self.sesion,
            medico=self.medico,
            nombre_paciente='Claudia Vanesa',
            apellido_paciente='Sanchez Barrios',
            dni_paciente='95951354',
            fecha_del_informe=registro_arterial.fecha_del_informe,
            tipo_obra_social='COBER',
            horario='EXTRA',
            cantidad_regiones=1,
            monto_calculado=Decimal('24200.00'),
        )
        RegistroEstudio.objects.create(
            registro=registro_venoso,
            estudio=self.venoso,
            cantidad=1,
        )

        resultado = auditar_cantidad_doppler_mmii(
            fecha_desde=date(2026, 6, 1),
            fecha_hasta=date(2026, 9, 30),
        )

        self.assertEqual(resultado['grupos_posible_duplicado'], 0)
        self.assertFalse(any(
            estudio.get('posible_duplicado')
            for registro in resultado['resultados']
            for estudio in registro['estudios']
        ))

    @override_settings(SECURE_SSL_REDIRECT=False)
    def test_pantalla_solo_lectura_permite_jefatura_y_no_residentes(self):
        registro = self._crear_registro()
        RegistroEstudio.objects.create(
            registro=registro,
            estudio=self.arterial,
            cantidad=2,
            contexto='SERVICIO',
        )
        RegistroEstudiosPorMedico.objects.filter(pk=registro.pk).update(
            cantidad_regiones=2,
            monto_calculado=Decimal('48400.00'),
        )
        url = reverse('liquidacion:auditoria_doppler_mmii')

        self.client.force_login(self.medico)
        response_residente = self.client.get(url)
        self.assertNotEqual(response_residente.status_code, 200)

        jefe = User.objects.create_user(
            username='jefatura_auditoria_doppler',
            first_name='Jefe',
            last_name='Servicio',
            rol='jefe_servicio',
            perfil_completo=True,
        )
        self.client.force_login(jefe)
        response_jefatura = self.client.get(url)

        self.assertEqual(response_jefatura.status_code, 200)
        self.assertContains(response_jefatura, 'Revisión sin ajustes económicos')
        self.assertContains(response_jefatura, 'Cantidad declarada')
        self.assertContains(response_jefatura, 'Impacto potencial')
        self.assertContains(response_jefatura, self.medico.get_full_name())

        response_filtrada = self.client.get(
            url,
            {'profesional': self.medico.pk},
        )
        self.assertEqual(response_filtrada.status_code, 200)
        self.assertContains(response_filtrada, 'Registros candidatos:')
        self.assertContains(response_filtrada, 'Ecodoppler arterial MM inferiores')

    @override_settings(SECURE_SSL_REDIRECT=False)
    def test_pantalla_muestra_posible_duplicado_y_referencias(self):
        primero = self._crear_registro()
        RegistroEstudio.objects.create(registro=primero, estudio=self.arterial, cantidad=1)
        segundo = RegistroEstudiosPorMedico.objects.create(
            sesion_contable=self.sesion,
            medico=self.medico,
            nombre_paciente=primero.nombre_paciente,
            apellido_paciente=primero.apellido_paciente,
            dni_paciente=primero.dni_paciente,
            fecha_del_informe=primero.fecha_del_informe,
            tipo_obra_social='COBER',
            horario='EXTRA',
            cantidad_regiones=1,
            monto_calculado=Decimal('24200.00'),
        )
        RegistroEstudio.objects.create(registro=segundo, estudio=self.arterial, cantidad=1)
        jefe = User.objects.create_user(
            username='jefatura_auditoria_duplicados',
            first_name='Jefe',
            last_name='Servicio',
            rol='jefe_servicio',
            perfil_completo=True,
        )
        self.client.force_login(jefe)

        response = self.client.get(reverse('liquidacion:auditoria_doppler_mmii'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Posible duplicado')
        self.assertContains(response, f'#{primero.pk}')
        self.assertContains(response, f'#{segundo.pk}')

    @override_settings(SECURE_SSL_REDIRECT=False)
    def test_pantalla_agrupa_art_y_venoso_en_un_registro_con_totales_separados(self):
        registro = self._crear_registro()
        RegistroEstudio.objects.create(
            registro=registro,
            estudio=self.arterial,
            cantidad=2,
            contexto='SERVICIO',
        )
        RegistroEstudio.objects.create(
            registro=registro,
            estudio=self.venoso,
            cantidad=2,
            contexto='SERVICIO',
        )
        RegistroEstudiosPorMedico.objects.filter(pk=registro.pk).update(
            cantidad_regiones=4,
            monto_calculado=Decimal('48400.00'),
        )
        jefe = User.objects.create_user(
            username='jefatura_auditoria_doppler_agrupada',
            first_name='Jefe',
            last_name='Servicio',
            rol='jefe_servicio',
            perfil_completo=True,
        )
        self.client.force_login(jefe)

        response = self.client.get(reverse('liquidacion:auditoria_doppler_mmii'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, f'Registro #{registro.pk}', count=1)
        self.assertContains(response, '4 declaradas')
        self.assertContains(response, '2 según regla')
        self.assertContains(response, '4 → 2')
        self.assertContains(response, 'Ecodoppler arterial MM inferiores')
        self.assertContains(response, 'Doppler venoso de miembros inferiores')
        self.assertContains(response, 'Cantidad declarada: <strong>2</strong>', count=2)

    def test_simular_cantidad_usa_calculo_canonico_sin_escribir(self):
        registro = self._crear_registro()
        relacion = RegistroEstudio.objects.create(
            registro=registro, estudio=self.arterial, cantidad=2, contexto='SERVICIO',
        )
        original = registro.calcular_monto()
        estimado = registro.calcular_monto(cantidades_auditoria={relacion.pk: 1})
        self.assertGreater(original, 0)
        self.assertEqual(estimado, original / 2)
        relacion.refresh_from_db()
        registro.refresh_from_db()
        self.assertEqual(relacion.cantidad, 2)
        self.assertEqual(registro.monto_calculado, Decimal('48400.00'))

    def _preparar_casos_economicos(self):
        registro = self._crear_registro()
        for estudio in (self.arterial, self.venoso):
            RegistroEstudio.objects.create(
                registro=registro, estudio=estudio, cantidad=2, contexto='SERVICIO',
            )
        RegistroEstudiosPorMedico.objects.filter(pk=registro.pk).update(
            monto_calculado=registro.calcular_monto(), cantidad_regiones=4,
        )
        registro.refresh_from_db()
        crear_casos_auditoria_doppler_mmii(auditar_cantidad_doppler_mmii(
            fecha_desde=date(2026, 6, 1), fecha_hasta=date(2026, 6, 30),
        ), self.medico)
        return registro, list(RevisionAuditoriaDopplerMMII.objects.order_by('pk'))

    def test_estimacion_total_suma_solo_confirmados_y_conserva_historial(self):
        registro, casos = self._preparar_casos_economicos()
        original = registro.monto_calculado
        for caso in casos:
            resolver_caso_auditoria_doppler_mmii(
                caso_id=caso.pk, decision='CONFIRMADO',
                evidencias={'orden_medica_verificada': True},
                observacion='Cantidad bilateral revisada', usuario=self.medico,
            )
        comparacion = adjuntar_comparacion_doppler_mmii([registro])[0].comparacion_doppler
        self.assertEqual(comparacion['monto_original'], original)
        self.assertEqual(comparacion['monto_estimado'], original / 2)
        self.assertEqual(comparacion['diferencia_estimada'], original / 2)
        self.assertFalse(comparacion['parcial'])
        primer_evento = HistorialRevisionAuditoriaDopplerMMII.objects.order_by('pk').first()
        self.assertEqual(Decimal(primer_evento.estimacion_json['monto_estimado']), original * Decimal('0.75'))
        resolver_caso_auditoria_doppler_mmii(
            caso_id=casos[0].pk, decision='DESCARTADO', evidencias={},
            observacion='Orden adicional justificada', usuario=self.medico,
        )
        comparacion = adjuntar_comparacion_doppler_mmii([registro])[0].comparacion_doppler
        self.assertEqual(comparacion['monto_estimado'], original * Decimal('0.75'))
        primer_evento.refresh_from_db()
        self.assertEqual(Decimal(primer_evento.estimacion_json['monto_estimado']), original * Decimal('0.75'))
        registro.refresh_from_db()
        self.assertEqual(registro.monto_calculado, original)
        self.assertEqual(list(registro.registroestudio_set.values_list('cantidad', flat=True)), [2, 2])

    def test_tarifa_no_reproduce_original_deja_estimacion_pendiente(self):
        registro, casos = self._preparar_casos_economicos()
        Estudios.objects.filter(pk=self.arterial.pk).update(precio_cober=Decimal('100.00'))
        resolver_caso_auditoria_doppler_mmii(
            caso_id=casos[0].pk, decision='CONFIRMADO',
            evidencias={'visualmedical_verificado': True},
            observacion='Diferencia de cantidad confirmada', usuario=self.medico,
        )
        comparacion = adjuntar_comparacion_doppler_mmii([registro])[0].comparacion_doppler
        self.assertIsNone(comparacion['monto_estimado'])
        self.assertIsNone(comparacion['diferencia_estimada'])
        self.assertIn('tarifa historica', comparacion['motivo'])

    def test_lote_confirma_solo_seleccionados_con_historial_y_sin_cambiar_montos(self):
        registro, casos = self._preparar_casos_economicos()
        jefe = User.objects.create_user(username='auditor_lote', rol='jefe_servicio', perfil_completo=True)
        confirmado = confirmar_lote_auditoria_doppler_mmii(
            caso_ids=[caso.pk for caso in casos], fecha_desde=date(2026, 6, 1),
            fecha_hasta=date(2026, 6, 30), medico_id=self.medico.pk,
            evidencias={'orden_medica_verificada': True},
            observacion='Se verificaron ordenes del lote bilateral', usuario=jefe,
        )
        self.assertEqual(confirmado, 2)
        self.assertEqual(HistorialRevisionAuditoriaDopplerMMII.objects.count(), 2)
        self.assertEqual(RevisionAuditoriaDopplerMMII.objects.filter(estado='CONFIRMADO').count(), 2)
        original = registro.monto_calculado
        registro.refresh_from_db()
        self.assertEqual(registro.monto_calculado, original)
        with self.assertRaises(ValidationError):
            confirmar_lote_auditoria_doppler_mmii(
                caso_ids=[casos[0].pk], fecha_desde=date(2026, 6, 1),
                fecha_hasta=date(2026, 6, 30), medico_id=None,
                evidencias={'orden_medica_verificada': True}, observacion='Reenvio', usuario=jefe,
            )
        self.assertEqual(HistorialRevisionAuditoriaDopplerMMII.objects.count(), 2)

    def test_lote_rechaza_seleccion_invalida_sin_decisiones_parciales(self):
        registro, casos = self._preparar_casos_economicos()
        jefe = User.objects.create_user(username='auditor_lote_invalido', rol='jefe_servicio')
        argumentos = {
            'caso_ids': [caso.pk for caso in casos], 'fecha_desde': date(2026, 6, 1),
            'fecha_hasta': date(2026, 6, 30), 'medico_id': None,
            'evidencias': {'eges_verificado': True}, 'observacion': 'Revision masiva', 'usuario': jefe,
        }
        with self.assertRaises(ValidationError):
            confirmar_lote_auditoria_doppler_mmii(**argumentos)
        self.assertFalse(HistorialRevisionAuditoriaDopplerMMII.objects.exists())
        argumentos['evidencias'] = {'orden_medica_verificada': True}
        fuente = dict(casos[1].datos_originales_json)
        fuente['posible_duplicado'] = True
        RevisionAuditoriaDopplerMMII.objects.filter(pk=casos[1].pk).update(datos_originales_json=fuente)
        with self.assertRaises(ValidationError):
            confirmar_lote_auditoria_doppler_mmii(**argumentos)
        self.assertFalse(HistorialRevisionAuditoriaDopplerMMII.objects.exists())
        self.assertEqual(RevisionAuditoriaDopplerMMII.objects.filter(estado='PENDIENTE').count(), 2)

    @override_settings(SECURE_SSL_REDIRECT=False)
    def test_endpoint_lote_respeta_filtros_permisos_y_evidencia(self):
        registro, casos = self._preparar_casos_economicos()
        url = reverse('liquidacion:auditoria_doppler_mmii_confirmar_lote')
        datos = {
            'casos': [caso.pk for caso in casos], 'fecha_desde': '2026-06-01',
            'fecha_hasta': '2026-06-30', 'profesional': self.medico.pk,
            'observacion': 'Regla bilateral verificada para todos los seleccionados',
        }
        self.client.force_login(self.medico)
        self.assertEqual(self.client.post(url, datos).status_code, 302)
        self.assertFalse(HistorialRevisionAuditoriaDopplerMMII.objects.exists())
        jefe = User.objects.create_user(username='auditor_endpoint_lote', rol='jefe_servicio', perfil_completo=True)
        self.client.force_login(jefe)
        response = self.client.get(reverse('liquidacion:auditoria_doppler_mmii'))
        self.assertContains(response, 'form="doppler-lote" data-doppler-caso', count=2)
        self.client.post(url, datos)
        self.assertFalse(HistorialRevisionAuditoriaDopplerMMII.objects.exists())
        self.client.post(url, {**datos, 'fecha_desde': '2026-02-31', 'orden_medica_verificada': 'on'})
        self.assertFalse(HistorialRevisionAuditoriaDopplerMMII.objects.exists())
        self.client.post(url, {**datos, 'fecha_hasta': '2026-06-05', 'orden_medica_verificada': 'on'})
        self.assertFalse(HistorialRevisionAuditoriaDopplerMMII.objects.exists())
        response = self.client.post(url, {**datos, 'orden_medica_verificada': 'on'})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(HistorialRevisionAuditoriaDopplerMMII.objects.count(), 2)

    @override_settings(SECURE_SSL_REDIRECT=False)
    def test_excel_personal_y_administrativo_comparten_estimacion_sin_cambiar_totales(self):
        registro, casos = self._preparar_casos_economicos()
        jefe = User.objects.create_user(username='auditor_excel_doppler', rol='jefe_servicio', perfil_completo=True)
        confirmar_lote_auditoria_doppler_mmii(
            caso_ids=[caso.pk for caso in casos], fecha_desde=date(2026, 6, 1),
            fecha_hasta=date(2026, 6, 30), medico_id=self.medico.pk,
            evidencias={'visualmedical_verificado': True}, observacion='=FUNDAMENTO_LITERAL', usuario=jefe,
        )
        self.client.force_login(self.medico)
        response = self.client.get(reverse('liquidacion:exportar_excel_mis_registros'), {'mes': 6, 'año': 2026})
        self.assertEqual(response.status_code, 200)
        personal = load_workbook(BytesIO(response.content))
        self.assertEqual(personal['Practicas']['J2'].value, float(registro.monto_calculado))
        self.assertEqual(personal['Practicas']['N2'].value, float(registro.monto_calculado / 2))
        self.assertEqual(personal['Practicas']['O2'].value, float(registro.monto_calculado / 2))
        self.assertEqual(personal['Resumen']['B5'].value, float(registro.monto_calculado))
        hoja = personal['Revision Doppler']
        self.assertEqual(hoja['H2'].value, '=FUNDAMENTO_LITERAL')
        self.assertEqual(hoja['H2'].data_type, 's')
        self.assertEqual(hoja['M2'].value, float(registro.monto_calculado / 2))
        self.assertIsNone(hoja['M3'].value)
        self.client.force_login(jefe)
        filtros = {'medico': self.medico.pk, 'mes': 6, 'año': 2026}
        response = self.client.get(reverse('liquidacion:exportar_excel_liquidacion'), filtros)
        self.assertEqual(response.status_code, 200)
        administrativo = load_workbook(BytesIO(response.content))
        principal = administrativo['Liquidación Completa']
        self.assertEqual(principal['I2'].value, float(registro.monto_calculado))
        self.assertEqual(principal['T2'].value, personal['Practicas']['N2'].value)
        self.assertEqual(principal['U2'].value, personal['Practicas']['O2'].value)
        self.assertEqual(administrativo['Auditoria Doppler']['R2'].value, hoja['M2'].value)
        self.sesion.estado = 'FACTURADA'
        self.sesion.save(update_fields=['estado'])
        response = self.client.get(reverse('liquidacion:exportar_excel_liquidacion_definitiva'), filtros)
        definitivo = load_workbook(BytesIO(response.content))
        self.assertEqual(definitivo['Liquidación Completa']['I2'].value, float(registro.monto_calculado))
        response = self.client.get(reverse('liquidacion:liquidacion_mensual'), filtros)
        self.assertContains(response, 'Comparación Doppler')
        self.assertContains(response, 'Débito Doppler: no aplicado.')
        self.assertEqual(response.context['medico_data'][0]['total_monto'], registro.monto_calculado)
        self.assertEqual(response.context['medico_data'][0]['resumen_doppler']['diferencia_estimada'], registro.monto_calculado / 2)

    def test_estimacion_no_se_recalcula_al_leer_y_se_oculta_si_fuentes_cambian(self):
        registro, casos = self._preparar_casos_economicos()
        resolver_caso_auditoria_doppler_mmii(
            caso_id=casos[0].pk, decision='CONFIRMADO',
            evidencias={'orden_medica_verificada': True}, observacion='Regla confirmada', usuario=self.medico,
        )
        monto_guardado = adjuntar_comparacion_doppler_mmii([registro])[0].comparacion_doppler['monto_estimado']
        Estudios.objects.filter(pk=self.arterial.pk).update(precio_cober=Decimal('1.00'))
        self.assertEqual(adjuntar_comparacion_doppler_mmii([registro])[0].comparacion_doppler['monto_estimado'], monto_guardado)
        RegistroEstudio.objects.filter(registro=registro, estudio=self.venoso).update(cantidad=3)
        self.assertIsNone(adjuntar_comparacion_doppler_mmii([registro])[0].comparacion_doppler['monto_estimado'])

    def test_residente_estimacion_informativa_no_se_suma_como_debito_propuesto(self):
        self.medico.rol = 'medico_residente'
        self.medico.save(update_fields=['rol'])
        registro = self._crear_registro(rol='medico_residente')
        RegistroEstudio.objects.create(registro=registro, estudio=self.arterial, cantidad=2, contexto='SERVICIO')
        RegistroEstudiosPorMedico.objects.filter(pk=registro.pk).update(monto_calculado=registro.calcular_monto())
        registro.refresh_from_db()
        crear_casos_auditoria_doppler_mmii(auditar_cantidad_doppler_mmii(
            fecha_desde=date(2026, 6, 1), fecha_hasta=date(2026, 6, 30),
        ), self.medico)
        caso = RevisionAuditoriaDopplerMMII.objects.get()
        resolver_caso_auditoria_doppler_mmii(
            caso_id=caso.pk, decision='CONFIRMADO', evidencias={'orden_medica_verificada': True},
            observacion='Revision informativa de residente', usuario=self.medico,
        )
        from .services_auditoria import resumir_comparacion_doppler_mmii
        comparados = adjuntar_comparacion_doppler_mmii([registro])
        self.assertTrue(comparados[0].comparacion_doppler['informativo_residente'])
        resumen = resumir_comparacion_doppler_mmii(comparados)
        self.assertEqual(resumen['diferencia_estimada'], Decimal('0.00'))
        self.assertEqual(resumen['diferencia_informativa_residentes'], registro.monto_calculado / 2)
        self.assertEqual(resumen['posible_debito'], 0)
        self.assertEqual(resumen['monto_proyectado_practicas'], registro.monto_calculado)
        self.assertEqual(valores_proyeccion_doppler_mmii(registro), [0, 0, float(registro.monto_calculado)])

    def test_resumen_proyectado_distingue_importes_pendientes_y_disponibles(self):
        registro, casos = self._preparar_casos_economicos()
        adjuntar_comparacion_doppler_mmii([registro])
        resumen = resumir_comparacion_doppler_mmii([registro])
        self.assertEqual(resumen['pendientes_economicos'], 1)
        self.assertEqual(valores_proyeccion_doppler_mmii(registro), ['Pendiente', 'Pendiente', 'Pendiente'])
        for caso in casos:
            resolver_caso_auditoria_doppler_mmii(
                caso_id=caso.pk, decision='CONFIRMADO', evidencias={'orden_medica_verificada': True},
                observacion='Orden revisada', usuario=self.medico,
            )
        adjuntar_comparacion_doppler_mmii([registro])
        resumen = resumir_comparacion_doppler_mmii([registro])
        self.assertEqual(resumen['pendientes_economicos'], 0)
        self.assertEqual(resumen['posible_debito'], registro.monto_calculado / 2)
        self.assertEqual(resumen['monto_proyectado_practicas'], registro.monto_calculado / 2)
        for caso in casos:
            resolver_caso_auditoria_doppler_mmii(
                caso_id=caso.pk, decision='DESCARTADO', evidencias={},
                observacion='Prestaciones justificadas', usuario=self.medico,
            )
        adjuntar_comparacion_doppler_mmii([registro])
        resumen = resumir_comparacion_doppler_mmii([registro])
        self.assertEqual(resumen['pendientes_economicos'], 0)
        self.assertEqual(resumen['posible_debito'], 0)
        self.assertEqual(valores_proyeccion_doppler_mmii(registro), [0, 0, float(registro.monto_calculado)])

    @override_settings(SECURE_SSL_REDIRECT=False)
    def test_excel_personal_incluye_periodo_completo_y_resume_posible_debito(self):
        registro, casos = self._preparar_casos_economicos()
        for caso in casos:
            resolver_caso_auditoria_doppler_mmii(
                caso_id=caso.pk, decision='CONFIRMADO', evidencias={'orden_medica_verificada': True},
                observacion='Orden bilateral revisada', usuario=self.medico,
            )
        for tipo, importe, dia in [('RES', Decimal('2000.00'), 12), ('ECO', Decimal('3000.00'), 13)]:
            estudio = Estudios.objects.create(
                nombre=f'Practica adicional {tipo}', tipo=tipo, precio_cober=importe,
                precio_otras_os=importe, conteo_regiones=1, conteo_regiones_default=1,
            )
            adicional = RegistroEstudiosPorMedico.objects.create(
                sesion_contable=self.sesion, medico=self.medico,
                nombre_paciente='Paciente', apellido_paciente=f'Prueba{tipo}',
                dni_paciente=f'123000{dia}', fecha_del_informe=date(2026, 6, dia),
                tipo_obra_social='COBER', horario='EXTRA', monto_calculado=importe,
            )
            RegistroEstudio.objects.create(registro=adicional, estudio=estudio, cantidad=1)
        guardia = GuardiaPasiva.objects.create(
            sesion_contable=self.sesion, medico=self.medico, fecha_guardia=date(2026, 6, 15),
            tipo_guardia='COBER',
        )
        self.client.force_login(self.medico)
        response = self.client.get(reverse('liquidacion:exportar_excel_mis_registros'), {
            'mes': 6, 'año': 2026, 'modalidad': 'DOP',
            'busqueda': 'NINGUN_PACIENTE_COINCIDE', 'filtro_rapido': 'hoy',
        })
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.active.title, 'Resumen')
        practicas = workbook['Practicas']
        self.assertEqual(practicas.max_row, 5)
        filas = list(practicas.iter_rows(min_row=2, max_row=4, values_only=True))
        self.assertEqual(sum(Decimal(str(fila[9])) for fila in filas), registro.monto_calculado + Decimal('5000.00'))
        self.assertTrue(any('Practica adicional RES' in fila[6] for fila in filas))
        self.assertTrue(any('Practica adicional ECO' in fila[6] for fila in filas))
        self.assertEqual(sum(Decimal(str(fila[13])) for fila in filas), registro.monto_calculado / 2)
        self.assertEqual(sum(Decimal(str(fila[14])) for fila in filas), registro.monto_calculado / 2 + Decimal('5000.00'))
        self.assertEqual(practicas['J5'].value, '=SUM(J2:J4)')
        self.assertEqual(practicas['N5'].value, float(registro.monto_calculado / 2))
        self.assertEqual(practicas['O5'].value, float(registro.monto_calculado / 2 + Decimal('5000.00')))
        resumen = workbook['Resumen']
        self.assertEqual(resumen['B7'].value, float(registro.monto_calculado + Decimal('5000.00') + guardia.monto))
        self.assertEqual(resumen['B9'].value, float(registro.monto_calculado / 2))
        self.assertEqual(resumen['B10'].value, float(registro.monto_calculado / 2 + Decimal('5000.00') + guardia.monto))
        self.assertIn('todas las modalidades', resumen['A14'].value)

    @override_settings(SECURE_SSL_REDIRECT=False)
    def test_excel_personal_compacto_distingue_residente_y_jefe(self):
        from .services import generar_buffer_excel_mis_registros

        registro, casos = self._preparar_casos_economicos()
        for caso in casos:
            resolver_caso_auditoria_doppler_mmii(
                caso_id=caso.pk, decision='CONFIRMADO', evidencias={'orden_medica_verificada': True},
                observacion='Regla verificada', usuario=self.medico,
            )
        def generar():
            return load_workbook(generar_buffer_excel_mis_registros(
                usuario=self.medico, mes=6, año=2026,
                registros=RegistroEstudiosPorMedico.objects.select_related('medico', 'sesion_contable').prefetch_related('registroestudio_set__estudio'),
                guardias=[],
            ))
        profesional = generar()
        self.assertEqual(profesional.active.title, 'Resumen')
        self.assertEqual(profesional['Resumen']['B7'].value, float(registro.monto_calculado))
        self.assertEqual(profesional['Resumen']['B9'].value, float(registro.monto_calculado / 2))
        self.assertEqual(profesional['Resumen']['B10'].value, float(registro.monto_calculado / 2))
        self.assertLessEqual(profesional['Resumen'].max_row, 15)
        self.medico.rol = 'medico_residente'
        self.medico.save(update_fields=['rol'])
        residente = generar()
        textos = ' '.join(str(celda.value or '') for hoja in residente for fila in hoja for celda in fila)
        self.assertNotIn('debito', textos.lower())
        self.assertNotIn('total estimado', textos.lower())
        self.assertEqual(residente['Practicas'].max_column, 13)
        self.assertEqual(residente['Revision Doppler'].max_column, 11)
        self.assertEqual(residente['Revision Doppler']['K1'].value, 'Fuentes verificadas')
        self.assertEqual(residente['Revision Doppler']['K2'].value, 'Orden medica')
        self.assertTrue(residente['Practicas'].column_dimensions['B'].hidden)
        self.assertEqual(residente['Resumen']['B7'].value, float(registro.monto_calculado))

    @override_settings(SECURE_SSL_REDIRECT=False)
    def test_excel_profesional_solo_muestra_credito_cuando_existe(self):
        registro = self._crear_registro()
        combinado = Estudios.objects.create(
            codigo='DOP-ART-VEN-EXPORT', nombre='Doppler arterial y venoso MMII', tipo='DOP',
            conteo_regiones=1, conteo_regiones_default=1, precio_cober=Decimal('1000.00'),
            precio_otras_os=Decimal('1000.00'),
        )
        RegistroEstudio.objects.create(registro=registro, estudio=combinado, cantidad=1)
        original = registro.calcular_monto()
        RegistroEstudiosPorMedico.objects.filter(pk=registro.pk).update(monto_calculado=original)
        crear_casos_auditoria_doppler_mmii(auditar_cantidad_doppler_mmii(
            fecha_desde=date(2026, 6, 1), fecha_hasta=date(2026, 6, 30),
        ), self.medico)
        caso = RevisionAuditoriaDopplerMMII.objects.get()
        resolver_caso_auditoria_doppler_mmii(
            caso_id=caso.pk, decision='CONFIRMADO', evidencias={'orden_medica_verificada': True},
            observacion='Dos modalidades documentadas', usuario=self.medico,
        )
        self.client.force_login(self.medico)
        response = self.client.get(reverse('liquidacion:exportar_excel_mis_registros'), {'mes': 6, 'año': 2026})
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook['Resumen']['A10'].value, 'Posible credito')
        self.assertEqual(workbook['Resumen']['B10'].value, float(original))
        self.assertEqual(workbook['Resumen']['B11'].value, float(original * 2))
        self.assertEqual(workbook['Practicas']['O1'].value, 'Posible credito (no aplicado)')
        self.assertEqual(workbook['Practicas']['P2'].value, float(original * 2))
        registro.refresh_from_db()
        self.assertEqual(registro.monto_calculado, original)

    @override_settings(SECURE_SSL_REDIRECT=False)
    def test_excel_personal_sin_registros_tiene_resumen_y_totales_validos(self):
        self.client.force_login(self.medico)
        response = self.client.get(reverse('liquidacion:exportar_excel_mis_registros'), {'mes': 6, 'año': 2026})
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.active.title, 'Resumen')
        self.assertEqual(workbook['Resumen']['B7'].value, 0)
        self.assertEqual(workbook['Resumen']['B9'].value, 0)
        self.assertEqual(workbook['Resumen']['B10'].value, 0)
        self.assertEqual(workbook['Practicas']['J2'].value, 0)
        self.assertEqual(workbook['Practicas'].auto_filter.ref, 'A1:O1')

    @override_settings(SECURE_SSL_REDIRECT=False)
    def test_excel_personal_muestra_pendientes_sin_inventar_importes(self):
        registro, casos = self._preparar_casos_economicos()
        self.client.force_login(self.medico)
        response = self.client.get(reverse('liquidacion:exportar_excel_mis_registros'), {'mes': 6, 'año': 2026})
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook['Practicas']['N2'].value, 'Pendiente')
        self.assertEqual(workbook['Practicas']['O2'].value, 'Pendiente')
        self.assertIn('Estimacion parcial', workbook['Resumen']['A11'].value)

    def test_persistencia_es_idempotente_y_no_modifica_fuentes_economicas(self):
        registro = self._crear_registro()
        relacion = RegistroEstudio.objects.create(
            registro=registro,
            estudio=self.arterial,
            cantidad=2,
            contexto='SERVICIO',
        )
        RegistroEstudiosPorMedico.objects.filter(pk=registro.pk).update(
            cantidad_regiones=4,
            monto_calculado=Decimal('48400.00'),
        )
        auditoria = auditar_cantidad_doppler_mmii(
            fecha_desde=date(2026, 6, 1),
            fecha_hasta=date(2026, 6, 30),
        )

        primera_generacion = crear_casos_auditoria_doppler_mmii(auditoria, self.medico)
        segunda_generacion = crear_casos_auditoria_doppler_mmii(auditoria, self.medico)

        self.assertEqual(primera_generacion, {'detectados': 1, 'creados': 1, 'existentes': 0})
        self.assertEqual(segunda_generacion, {'detectados': 1, 'creados': 0, 'existentes': 1})
        self.assertEqual(RevisionAuditoriaDopplerMMII.objects.count(), 1)
        caso = RevisionAuditoriaDopplerMMII.objects.get()
        self.assertEqual(caso.datos_originales_json['cantidad_declarada'], 2)
        self.assertEqual(caso.datos_originales_json['cantidad_esperada'], 1)
        self.assertEqual(caso.datos_originales_json['monto_registrado'], '48400.00')

        registro.refresh_from_db()
        relacion.refresh_from_db()
        self.assertEqual(registro.cantidad_regiones, 4)
        self.assertEqual(registro.monto_calculado, Decimal('48400.00'))
        self.assertEqual(relacion.cantidad, 2)

    def test_confirmacion_exige_evidencia_y_cada_decision_agrega_historial(self):
        registro = self._crear_registro()
        RegistroEstudio.objects.create(
            registro=registro,
            estudio=self.arterial,
            cantidad=2,
            contexto='SERVICIO',
        )
        auditoria = auditar_cantidad_doppler_mmii(
            fecha_desde=date(2026, 6, 1),
            fecha_hasta=date(2026, 6, 30),
        )
        crear_casos_auditoria_doppler_mmii(auditoria, self.medico)
        caso = RevisionAuditoriaDopplerMMII.objects.get()

        with self.assertRaises(ValidationError):
            resolver_caso_auditoria_doppler_mmii(
                caso_id=caso.pk,
                decision=RevisionAuditoriaDopplerMMII.ESTADO_CONFIRMADO,
                evidencias={},
                observacion='Sin evidencia documental',
                usuario=self.medico,
            )

        resolver_caso_auditoria_doppler_mmii(
            caso_id=caso.pk,
            decision=RevisionAuditoriaDopplerMMII.ESTADO_CONFIRMADO,
            evidencias={'orden_medica_verificada': True},
            observacion='Orden médica revisada',
            usuario=self.medico,
        )
        resolver_caso_auditoria_doppler_mmii(
            caso_id=caso.pk,
            decision=RevisionAuditoriaDopplerMMII.ESTADO_DESCARTADO,
            evidencias={'eges_verificado': True},
            observacion='Se descarta tras revisión complementaria',
            usuario=self.medico,
        )

        caso.refresh_from_db()
        self.assertEqual(caso.estado, RevisionAuditoriaDopplerMMII.ESTADO_DESCARTADO)
        self.assertEqual(caso.observacion, 'Se descarta tras revisión complementaria')
        self.assertEqual(HistorialRevisionAuditoriaDopplerMMII.objects.count(), 2)
        self.assertEqual(
            list(caso.historial.order_by('pk').values_list('estado_anterior', 'estado_nuevo')),
            [
                (RevisionAuditoriaDopplerMMII.ESTADO_PENDIENTE, RevisionAuditoriaDopplerMMII.ESTADO_CONFIRMADO),
                (RevisionAuditoriaDopplerMMII.ESTADO_CONFIRMADO, RevisionAuditoriaDopplerMMII.ESTADO_DESCARTADO),
            ],
        )
        registro.refresh_from_db()
        self.assertEqual(registro.monto_calculado, Decimal('48400.00'))

    @override_settings(SECURE_SSL_REDIRECT=False)
    def test_endpoints_persisten_casos_y_restringen_acceso_a_jefatura(self):
        registro = self._crear_registro()
        RegistroEstudio.objects.create(
            registro=registro,
            estudio=self.arterial,
            cantidad=2,
            contexto='SERVICIO',
        )
        generar_url = reverse('liquidacion:auditoria_doppler_mmii_generar_casos')
        datos = {
            'fecha_desde': '2026-06-01',
            'fecha_hasta': '2026-06-30',
            'profesional': '',
            'solo_diferencias': '1',
        }

        self.client.force_login(self.medico)
        respuesta_no_autorizada = self.client.post(generar_url, datos)
        self.assertEqual(respuesta_no_autorizada.status_code, 302)
        self.assertEqual(RevisionAuditoriaDopplerMMII.objects.count(), 0)

        jefe = User.objects.create_user(
            username='jefatura_generacion_doppler',
            rol='jefe_servicio',
            perfil_completo=True,
        )
        self.client.force_login(jefe)
        respuesta_fecha_invalida = self.client.post(generar_url, {
            **datos,
            'fecha_desde': '2026-02-31',
        })
        self.assertEqual(respuesta_fecha_invalida.status_code, 302)
        self.assertEqual(RevisionAuditoriaDopplerMMII.objects.count(), 0)
        respuesta = self.client.post(generar_url, datos)
        self.assertEqual(respuesta.status_code, 302)
        self.assertEqual(RevisionAuditoriaDopplerMMII.objects.count(), 1)

        caso = RevisionAuditoriaDopplerMMII.objects.get()
        resolver_url = reverse(
            'liquidacion:auditoria_doppler_mmii_resolver',
            args=[caso.pk],
        )
        self.client.force_login(self.medico)
        self.client.post(resolver_url, {
            **datos,
            'decision': RevisionAuditoriaDopplerMMII.ESTADO_DESCARTADO,
            'observacion': 'Intento sin permisos',
        })
        self.assertEqual(HistorialRevisionAuditoriaDopplerMMII.objects.count(), 0)
        self.client.force_login(jefe)
        self.client.post(resolver_url, {
            **datos,
            'decision': RevisionAuditoriaDopplerMMII.ESTADO_CONFIRMADO,
            'eges_verificado': 'on',
            'observacion': 'EGES solo no confirma',
        })
        self.assertEqual(HistorialRevisionAuditoriaDopplerMMII.objects.count(), 0)
        respuesta_revision = self.client.post(resolver_url, {
            **datos,
            'decision': RevisionAuditoriaDopplerMMII.ESTADO_CONFIRMADO,
            'orden_medica_verificada': 'on',
            'observacion': 'Orden verificada en prueba',
        })
        self.assertEqual(respuesta_revision.status_code, 302)
        self.assertEqual(HistorialRevisionAuditoriaDopplerMMII.objects.count(), 1)

    @override_settings(SECURE_SSL_REDIRECT=False)
    def test_mis_registros_resume_doppler_sin_indicar_sin_revision(self):
        registro, casos = self._preparar_casos_economicos()
        self.client.force_login(self.medico)
        response = self.client.get(
            reverse('liquidacion:registroestudios_list'), {'mes': 6, 'año': 2026},
        )
        self.assertContains(response, 'data-revision-resumen', count=1)
        self.assertContains(response, 'data-revision-detalle', count=1)
        self.assertContains(response, f'aria-controls="detalle-revision-{registro.pk}"', count=1)
        self.assertContains(response, 'aria-expanded="false"', count=1)
        self.assertContains(response, 'Auditoría Doppler MMII')
        self.assertNotContains(response, 'Sin revisión')

    @override_settings(SECURE_SSL_REDIRECT=False)
    def test_revision_horaria_y_doppler_comparten_un_detalle(self):
        registro, casos = self._preparar_casos_economicos()
        SolicitudRevisionHorarioRegistro.objects.create(
            registro=registro, solicitado_por=self.medico, horario_solicitado='EXTRA',
            fecha_hora_real_declarada=timezone.now(), motivo_solicitud='MOTIVO_HORARIO_PROPIO',
        )
        self.client.force_login(self.medico)
        response = self.client.get(
            reverse('liquidacion:registroestudios_list'), {'mes': 6, 'año': 2026},
        )
        self.assertContains(response, f'id="detalle-revision-{registro.pk}"', count=1)
        self.assertContains(response, 'data-revision-detalle', count=1)
        self.assertContains(response, 'MOTIVO_HORARIO_PROPIO')
        self.assertContains(response, 'Revisión pendiente')
        self.assertContains(response, 'Auditoría Doppler MMII')
        self.assertNotContains(response, 'Sin revisión')

    @override_settings(SECURE_SSL_REDIRECT=False)
    def test_mis_registros_solo_muestra_auditorias_del_profesional_autenticado(self):
        registro_propio = self._crear_registro()
        relacion_propia = RegistroEstudio.objects.create(
            registro=registro_propio,
            estudio=self.arterial,
            cantidad=2,
            contexto='SERVICIO',
        )
        otro_medico = User.objects.create_user(
            username='otro_medico_auditoria',
            rol='jefe_residentes',
        )
        otro_registro = RegistroEstudiosPorMedico.objects.create(
            sesion_contable=self.sesion,
            medico=otro_medico,
            nombre_paciente='Paciente',
            apellido_paciente='Reservado',
            dni_paciente='12345678',
            fecha_del_informe=date(2026, 6, 11),
            tipo_obra_social='COBER',
            horario='EXTRA',
            cantidad_regiones=2,
            monto_calculado=Decimal('24200.00'),
        )
        relacion_ajena = RegistroEstudio.objects.create(
            registro=otro_registro,
            estudio=self.arterial,
            cantidad=2,
            contexto='SERVICIO',
        )
        auditoria = auditar_cantidad_doppler_mmii(
            fecha_desde=date(2026, 6, 1),
            fecha_hasta=date(2026, 6, 30),
        )
        crear_casos_auditoria_doppler_mmii(auditoria, self.medico)
        casos = {
            caso.registro_estudio_id_origen: caso
            for caso in RevisionAuditoriaDopplerMMII.objects.all()
        }
        resolver_caso_auditoria_doppler_mmii(
            caso_id=casos[relacion_propia.pk].pk,
            decision=RevisionAuditoriaDopplerMMII.ESTADO_DESCARTADO,
            evidencias={},
            observacion='OBSERVACION_PROPIA_UNICA',
            usuario=self.medico,
        )
        resolver_caso_auditoria_doppler_mmii(
            caso_id=casos[relacion_ajena.pk].pk,
            decision=RevisionAuditoriaDopplerMMII.ESTADO_DESCARTADO,
            evidencias={},
            observacion='OBSERVACION_AJENA_NO_VISIBLE',
            usuario=self.medico,
        )

        self.client.force_login(self.medico)
        response = self.client.get(
            reverse('liquidacion:registroestudios_list'),
            {'mes': 6, 'año': 2026},
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'OBSERVACION_PROPIA_UNICA')
        self.assertNotContains(response, 'OBSERVACION_AJENA_NO_VISIBLE')
        self.assertNotContains(response, 'Paciente Reservado')
